"""
rto_reports.py - New-car reports from the delivery (RTO) list, and the
executive summary (v0.9.0)
=======================================================================

WHAT THIS MODULE ADDS TO THE WORKBOOK
-------------------------------------
    Summary                    one-page executive summary (always written)
  and, when the dealership's delivery list is given on Generate (step 4):
    13 New-car penetration     cars delivered vs cars that took DNS
                               accessories, OE vs DNS value - by location
                               and by model
    14 Missed opportunity      delivered cars WITHOUT DNS accessories, with
                               the remark given, and a count by remark
    15 RTO list vs invoices    each car's DNS value in the list against what
                               was actually invoiced - differences flagged
    16 New-car vs other        sales and profit from delivered cars against
                               all other business; profit per car by model
                               and by location
    17 Consultant scorecard    per sales consultant: cars delivered, cars
                               converted, value, profit, and the incentive
                               in the list next to the tool's calculation

HOW A DELIVERED CAR IS LINKED TO INVOICES
-----------------------------------------
By the last six digits of the VIN (see app/data/rto_list.py). An invoice
whose VIN field / customer name ends in six digits that are in the list
belongs to that car. A car can have more than one invoice. Invoices that
link to no delivered car are "other business" (walk-in, older cars,
counter sales).

WHAT THE FIGURES MEAN
---------------------
    "as per list"   the dealership's own columns (DNS Total Value, Executive
                    Incentive) - what the staff recorded
    "invoiced"      the tool's figures from Zoho for the linked invoices
The two are shown side by side and NOT forced to agree: the differences
are exactly what report 15 is for. A car "took DNS" for penetration when
the list shows a DNS value or list, OR an invoice is linked to it.

Sales are after discount and exclude GST (as everywhere in the workbook);
"Invoice total" includes GST, to compare with the list's DNS value.

No database and no Qt code here: MonthData + the list in, sheets out.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime

from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.data.rto_list import RtoCar, vin6
from app.reports.data import Invoice, MonthData
from app.reports.workbook import (
    AMBER_TINT, MONEY, NAVY, PCT, SLATE, Col, _finish, _font, exec_label,
    section, table, title)

RTO_SHEETS = ("13 New-car penetration", "14 Missed opportunity",
              "15 RTO list vs invoices", "16 New-car vs other",
              "17 Consultant scorecard")
SUMMARY_SHEET = "Summary"
AGREE_WITHIN = 1.0          # rupees: list value vs invoice total "agrees"


# ---------------------------------------------------------------------------
# Linking the list to the invoices
# ---------------------------------------------------------------------------
class Linked:
    """The delivery list joined to the month's invoices."""

    def __init__(self, d: MonthData, cars: list[RtoCar]):
        self.cars = cars
        by_vin = {c.vin6: c for c in cars if c.vin6}
        self.invoices_of: dict[int, list[Invoice]] = defaultdict(list)  # car row -> invoices
        self.other: list[Invoice] = []
        for inv in d.invoices:
            car = by_vin.get(vin6(inv.vin) or vin6(inv.customer))
            if car is not None:
                self.invoices_of[car.row].append(inv)
            else:
                self.other.append(inv)

    def invoices(self, car: RtoCar) -> list[Invoice]:
        return self.invoices_of.get(car.row, [])

    def took_dns(self, car: RtoCar) -> bool:
        return car.has_dns or bool(self.invoices(car))

    @property
    def linked(self) -> list[Invoice]:
        return [i for invs in self.invoices_of.values() for i in invs]


def _sum(invoices: list[Invoice], attr: str) -> float:
    return round(sum(getattr(i, attr) for i in invoices), 2)


def _ratio(a: float, b: float):
    return a / b if b else ""


# ---------------------------------------------------------------------------
# 13 New-car penetration
# ---------------------------------------------------------------------------
def _group_rows(link: Linked, key) -> list[dict]:
    groups: dict[str, dict] = {}
    for c in link.cars:
        g = groups.setdefault(key(c), dict(name=key(c), cars=0, dns=0, dns_value=0.0,
                                           oe=0, oe_value=0.0, oe_blank=0, none=0,
                                           invoiced=0.0))
        g["cars"] += 1
        g["dns"] += link.took_dns(c)
        g["dns_value"] += c.dns_value
        g["oe"] += c.has_oe
        g["oe_value"] += c.oe_value
        g["oe_blank"] += bool(c.oe_list) and not c.oe_value
        g["none"] += not link.took_dns(c) and not c.has_oe
        g["invoiced"] += _sum(link.invoices(c), "total")
    rows = sorted(groups.values(), key=lambda g: -g["cars"])
    for g in rows:
        for k in ("dns_value", "oe_value", "invoiced"):
            g[k] = round(g[k], 2)
    return rows


def _penetration_cols(first: str) -> list[Col]:
    pen = '=IF({cars}{r}=0,"",{dns}{r}/{cars}{r})'
    per = '=IF({cars}{r}=0,"",{dns_value}{r}/{cars}{r})'
    return [
        Col(first, "name", width=24),
        Col("Cars delivered", "cars", "qty", 10, total="sum"),
        Col("Took DNS", "dns", "qty", 10, total="sum"),
        Col("Penetration %", "pen", "pct", 11, formula=pen, total=pen),
        Col("DNS value (list)", "dns_value", "money", 14, total="sum"),
        Col("DNS value per car delivered", "per", "money", 14, formula=per, total=per),
        Col("Invoiced (with GST)", "invoiced", "money", 14, total="sum"),
        Col("Took OE", "oe", "qty", 9, total="sum"),
        Col("OE value (list)", "oe_value", "money", 14, total="sum"),
        Col("OE listed, value blank", "oe_blank", "qty", 11, total="sum"),
        Col("Neither OE nor DNS", "none", "qty", 10, total="sum"),
    ]


def penetration_sheet(ws: Worksheet, d: MonthData, link: Linked) -> None:
    row = title(ws, "New-car penetration", d.label, [
        "Source: the dealership's delivery (RTO) list. A car 'took DNS' when the list shows "
        "DNS accessories or an invoice is linked to it (last six digits of the VIN).",
        "OE = the car maker's accessories. 'OE listed, value blank' = accessories are named "
        "in the list but no value was filled in, so OE value is understated."])
    row = section(ws, row, "By location")
    t = table(ws, row, _penetration_cols("Location"), _group_rows(link, lambda c: c.location))
    row = section(ws, t["next"], "By model")
    table(ws, row, _penetration_cols("Model"), _group_rows(link, lambda c: c.model_group))
    _finish(ws)


# ---------------------------------------------------------------------------
# 14 Missed opportunity
# ---------------------------------------------------------------------------
def missed_sheet(ws: Worksheet, d: MonthData, link: Linked) -> None:
    missed = [c for c in link.cars if not link.took_dns(c)]
    row = title(ws, "Missed opportunity", d.label, [
        "Cars delivered this month with no DNS accessories (none in the list and no invoice "
        "linked). Remarks are as written in the list.",
        "Remarks such as 'will update later' or 'waiting for customer' are a follow-up list."])
    counts: dict[str, int] = defaultdict(int)
    for c in missed:
        counts[c.remarks.upper() or "(no remark)"] += 1
    row = section(ws, row, "By remark")
    t = table(ws, row, [
        Col("Remark", "remark", width=40),
        Col("Cars", "cars", "qty", 10, total="sum"),
    ], [dict(remark=k, cars=v) for k, v in sorted(counts.items(), key=lambda x: -x[1])],
        empty_text="Every delivered car took DNS accessories.")
    row = section(ws, t["next"], "Cars without DNS accessories")
    table(ws, row, [
        Col("Delivery date", "date", "date", 12),
        Col("Customer", "customer", width=28),
        Col("Model", "model", width=16),
        Col("VIN", "vin", width=20),
        Col("Sales consultant", "consultant", width=24),
        Col("Location", "location", width=12),
        Col("OE accessories", "oe_list", width=36),
        Col("OE value", "oe_value", "money", 12, total="sum"),
        Col("Remark", "remarks", width=32),
    ], [dict(date=c.delivered, customer=c.customer, model=c.model, vin=c.vin,
             consultant=c.consultant, location=c.location, oe_list=c.oe_list,
             oe_value=c.oe_value, remarks=c.remarks) for c in missed],
        filters=True, empty_text="Every delivered car took DNS accessories.")
    _finish(ws)


# ---------------------------------------------------------------------------
# 15 RTO list vs invoices
# ---------------------------------------------------------------------------
def _status(car: RtoCar, invoices: list[Invoice]) -> str:
    total = _sum(invoices, "total")
    if not invoices:
        return "In the list, no invoice found"
    if car.dns_value <= 0:
        return "Invoiced, list shows no DNS value"
    return "Agrees" if abs(total - car.dns_value) <= AGREE_WITHIN else "Amount differs"


def reconcile_rows(link: Linked) -> list[dict]:
    rows = []
    for c in link.cars:
        invs = link.invoices(c)
        if c.dns_value <= 0 and not invs:
            continue
        rows.append(dict(
            date=c.delivered, customer=c.customer, model=c.model, vin=c.vin,
            consultant=c.consultant, location=c.location, dns_list=c.dns_list,
            listed=c.dns_value, invoices=", ".join(i.invoice_no for i in invs),
            invoiced=_sum(invs, "total"), status=_status(c, invs)))
    order = {"In the list, no invoice found": 0, "Invoiced, list shows no DNS value": 1,
             "Amount differs": 2, "Agrees": 3}
    return sorted(rows, key=lambda r: (order[r["status"]], r["customer"]))


def reconcile_sheet(ws: Worksheet, d: MonthData, link: Linked) -> None:
    rows = reconcile_rows(link)
    row = title(ws, "RTO list vs invoices", d.label, [
        "Each delivered car with a DNS value in the list or an invoice in Zoho. 'DNS value "
        "(list)' is the dealership's figure; 'Invoiced' is the total of the linked invoices, "
        "including GST.",
        "'In the list, no invoice found': not billed this month, billed without the VIN "
        "digits, or left out of the reports (see 'Not included'). Differences first."])
    counts: dict[str, dict] = {}
    for r in rows:
        c = counts.setdefault(r["status"], dict(status=r["status"], cars=0, listed=0.0,
                                                invoiced=0.0))
        c["cars"] += 1
        c["listed"] = round(c["listed"] + r["listed"], 2)
        c["invoiced"] = round(c["invoiced"] + r["invoiced"], 2)
    row = section(ws, row, "Summary")
    t = table(ws, row, [
        Col("Result", "status", width=34),
        Col("Cars", "cars", "qty", 8, total="sum"),
        Col("DNS value (list)", "listed", "money", 15, total="sum"),
        Col("Invoiced (with GST)", "invoiced", "money", 15, total="sum"),
        Col("Difference", "diff", "money", 14, formula="={invoiced}{r}-{listed}{r}",
            total="sum"),
    ], list(counts.values()), empty_text="No car in the list has DNS accessories.")
    row = section(ws, t["next"], "Car by car")
    table(ws, row, [
        Col("Result", "status", width=34),
        Col("Delivery date", "date", "date", 12),
        Col("Customer", "customer", width=28),
        Col("Model", "model", width=15),
        Col("VIN", "vin", width=20),
        Col("Sales consultant", "consultant", width=22),
        Col("Location", "location", width=10),
        Col("DNS accessories (list)", "dns_list", width=34),
        Col("DNS value (list)", "listed", "money", 14, total="sum"),
        Col("Invoice no", "invoices", width=24),
        Col("Invoiced (with GST)", "invoiced", "money", 14, total="sum"),
        Col("Difference", "diff", "money", 13, formula="={invoiced}{r}-{listed}{r}",
            total="sum"),
    ], rows, filters=True, empty_text="No car in the list has DNS accessories.")
    _finish(ws)


# ---------------------------------------------------------------------------
# 16 New-car vs other business
# ---------------------------------------------------------------------------
GP = "={sales}{r}-{cost}{r}-{labour}{r}"
MARGIN = '=IF({sales}{r}=0,"",{gp}{r}/{sales}{r})'


def _business(name: str, invoices: list[Invoice]) -> dict:
    return dict(name=name, invoices=len(invoices), sales=_sum(invoices, "sales"),
                cost=_sum(invoices, "cost"), labour=_sum(invoices, "labour"))


def source_sheet(ws: Worksheet, d: MonthData, link: Linked) -> None:
    row = title(ws, "New-car vs other business", d.label, [
        "New-car business = invoices linked to a car in this month's delivery list. Other = "
        "every other invoice (walk-in, earlier deliveries, counter sales).",
        "Sales exclude GST. Gross profit = sales - product cost - labour."])
    row = section(ws, row, "Where the month's business came from")
    avg = '=IF({invoices}{r}=0,"",{sales}{r}/{invoices}{r})'
    t = table(ws, row, [
        Col("Business", "name", width=30),
        Col("Invoices", "invoices", "qty", 10, total="sum"),
        Col("Sales", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 10, formula=MARGIN, total=MARGIN),
        Col("Average bill", "avg", "money", 13, formula=avg, total=avg),
    ], [_business("New cars delivered this month", link.linked),
        _business("Other business", link.other)])

    def per_car(key) -> list[dict]:
        groups: dict[str, dict] = {}
        for c in link.cars:
            g = groups.setdefault(key(c), dict(name=key(c), cars=0, invoiced=0, sales=0.0,
                                               cost=0.0, labour=0.0))
            invs = link.invoices(c)
            g["cars"] += 1
            g["invoiced"] += bool(invs)
            for k in ("sales", "cost", "labour"):
                g[k] = round(g[k] + _sum(invs, k), 2)
        return sorted(groups.values(), key=lambda g: -g["sales"])

    def cols(first: str) -> list[Col]:
        return [
            Col(first, "name", width=30),
            Col("Cars delivered", "cars", "qty", 10, total="sum"),
            Col("Cars invoiced", "invoiced", "qty", 10, total="sum"),
            Col("Sales", "sales", "money", 14, total="sum"),
            Col("Product cost", "cost", "money", 14, total="sum"),
            Col("Labour", "labour", "money", 12, total="sum"),
            Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
            Col("Margin %", "margin", "pct", 10, formula=MARGIN, total=MARGIN),
            Col("Profit per car invoiced", "p1", "money", 13,
                formula='=IF({invoiced}{r}=0,"",{gp}{r}/{invoiced}{r})',
                total='=IF({invoiced}{r}=0,"",{gp}{r}/{invoiced}{r})'),
            Col("Profit per car delivered", "p2", "money", 13,
                formula='=IF({cars}{r}=0,"",{gp}{r}/{cars}{r})',
                total='=IF({cars}{r}=0,"",{gp}{r}/{cars}{r})'),
        ]
    row = section(ws, t["next"], "Profit per delivered car - by model")
    t = table(ws, row, cols("Model"), per_car(lambda c: c.model_group))
    row = section(ws, t["next"], "Profit per delivered car - by location")
    table(ws, row, cols("Location"), per_car(lambda c: c.location))
    _finish(ws)


# ---------------------------------------------------------------------------
# 17 Consultant scorecard
# ---------------------------------------------------------------------------
def scorecard_sheet(ws: Worksheet, d: MonthData, link: Linked) -> None:
    row = title(ws, "Consultant scorecard", d.label, [
        "Sales consultants as named in the delivery list. 'Incentive (list)' is the Executive "
        "Incentive column of the list; 'Incentive (tool)' is the spot incentive the tool "
        "works out for the invoices linked to the consultant's cars.",
        "The two are shown side by side for checking; neither is changed."])
    groups: dict[tuple, dict] = {}
    for c in link.cars:
        key = (c.consultant.upper() or "(not given)", c.location)
        g = groups.setdefault(key, dict(consultant=key[0], location=key[1], cars=0, dns=0,
                                        listed=0.0, sales=0.0, cost=0.0, labour=0.0,
                                        inc_list=0.0, inc_tool=0.0))
        invs = link.invoices(c)
        g["cars"] += 1
        g["dns"] += link.took_dns(c)
        g["listed"] += c.dns_value
        g["inc_list"] += c.incentive
        for k in ("sales", "cost", "labour"):
            g[k] += _sum(invs, k)
        g["inc_tool"] += sum(i.incentive_payable for i in invs)
    rows = sorted(groups.values(), key=lambda g: (-g["sales"], -g["cars"]))
    for g in rows:
        for k in ("listed", "sales", "cost", "labour", "inc_list", "inc_tool"):
            g[k] = round(g[k], 2)
    pen = '=IF({cars}{r}=0,"",{dns}{r}/{cars}{r})'
    table(ws, row, [
        Col("Sales consultant", "consultant", width=28),
        Col("Location", "location", width=12),
        Col("Cars delivered", "cars", "qty", 10, total="sum"),
        Col("Took DNS", "dns", "qty", 9, total="sum"),
        Col("Penetration %", "pen", "pct", 11, formula=pen, total=pen),
        Col("DNS value (list)", "listed", "money", 14, total="sum"),
        Col("Sales invoiced", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 13, total="sum"),
        Col("Labour", "labour", "money", 11, total="sum"),
        Col("Gross profit", "gp", "money", 13, formula=GP, total="sum"),
        Col("Incentive (list)", "inc_list", "money", 13, total="sum"),
        Col("Incentive (tool)", "inc_tool", "money", 13, total="sum"),
        Col("Difference", "diff", "money", 12, formula="={inc_tool}{r}-{inc_list}{r}",
            total="sum"),
    ], rows, filters=True)
    _finish(ws, "C" + str(row + 1))


# ---------------------------------------------------------------------------
# Executive summary
# ---------------------------------------------------------------------------
def summary_sheet(ws: Worksheet, d: MonthData, previous: MonthData | None,
                  link: Linked | None, user: str = "", attention: bool = True,
                  highest_first: bool = False) -> None:
    """
    One page for the owner: the month's headline figures (with last month's
    where a previous month has been read), the new-car business, the best
    and weakest performers, packages, costs, and points needing attention.
    Figures are values (not formulas) so the page reads the same anywhere.
    """
    ws["A1"] = "Drive N Style – Executive summary"
    ws["A1"].font = _font(True, NAVY, 16)
    ws["A2"] = d.label
    ws["A2"].font = _font(True, SLATE, 12)
    ws["A3"] = (f"Prepared on {datetime.now():%d-%m-%Y %H:%M}"
                + (f" by {user}" if user else ""))
    ws["A3"].font = _font(color=SLATE, size=9)
    state = {"row": 5}

    def head(text: str, *cols: str) -> None:
        r = state["row"] = state["row"] + 1
        for c, t in enumerate((text, *cols), start=1):
            cell = ws.cell(r, c, t)
            cell.font = _font(True, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "right")
        state["row"] += 1

    def put(label: str, *values, fmt=MONEY, bold=False) -> None:
        r = state["row"]
        ws.cell(r, 1, label).font = _font(bold)
        for c, v in enumerate(values, start=2):
            cell = ws.cell(r, c, v)
            cell.font = _font(bold)
            cell.alignment = Alignment(horizontal="right")
            if isinstance(v, (int, float)):
                cell.number_format = fmt[c - 2] if isinstance(fmt, tuple) else fmt
        state["row"] += 1

    def figures(m: MonthData) -> dict:
        indirect = round(sum(a for _, a in m.entered_indirect)
                         + sum(a for _, _, a in m.auto_indirect), 2)
        n = len(m.invoices)
        return dict(sales=m.sales, gp=m.gross_profit,
                    margin=_ratio(m.gross_profit, m.sales), indirect=indirect,
                    net=round(m.gross_profit - indirect, 2), invoices=n,
                    avg=round(m.sales / n, 2) if n else 0.0, cogs=m.cogs)

    now, before = figures(d), figures(previous) if previous else None
    cols = ["This month"] + ([previous.label, "Change"] if previous else [])
    head("Headline figures", *cols)
    for label, key, fmt in (("Sales (excluding GST)", "sales", MONEY),
                            ("Gross profit", "gp", MONEY),
                            ("Gross margin %", "margin", PCT),
                            ("Indirect costs", "indirect", MONEY),
                            ("Net profit", "net", MONEY),
                            ("Invoices", "invoices", "0"),
                            ("Average bill", "avg", MONEY)):
        values = [now[key]]
        if before:
            values += [before[key],
                       now[key] - before[key] if isinstance(before[key], (int, float))
                       and isinstance(now[key], (int, float)) else ""]
        put(label, *values, fmt=fmt, bold=key in ("sales", "net"))

    show_attention, attention = attention, []
    if link is not None:
        cars = len(link.cars)
        took = sum(link.took_dns(c) for c in link.cars)
        listed = round(sum(c.dns_value for c in link.cars), 2)
        oe = round(sum(c.oe_value for c in link.cars), 2)
        new_sales, other_sales = _sum(link.linked, "sales"), _sum(link.other, "sales")
        head("New-car business (delivery list)", "")
        put("Cars delivered", cars, fmt="0")
        put("Cars that took DNS accessories", took, fmt="0")
        put("DNS penetration %", _ratio(took, cars), fmt=PCT, bold=True)
        put("DNS value as per list", listed)
        put("DNS value per car delivered", round(listed / cars, 2) if cars else 0.0)
        put("OE accessories value as per list", oe)
        put("Sales from delivered cars (excl. GST)", new_sales)
        put("Sales from other business (excl. GST)", other_sales)
        put("Share of sales from delivered cars", _ratio(new_sales, d.sales), fmt=PCT)

        by_loc = _group_rows(link, lambda c: c.location)
        if highest_first:     # PDF version (v0.10.5): highest penetration % first
            by_loc.sort(key=lambda g: -(g["dns"] / g["cars"] if g["cars"] else 0))
        head("Penetration by location", "Cars delivered", "Took DNS", "Penetration %")
        for g in by_loc:
            put(g["name"], g["cars"], g["dns"], _ratio(g["dns"], g["cars"]),
                fmt=("0", "0", PCT))
        weak = [g for g in by_loc if g["cars"] >= 5 and g["name"] != "(not given)"]
        if weak:
            w = min(weak, key=lambda g: g["dns"] / g["cars"])
            attention.append(f"{w['name']}: {w['cars']} cars delivered, only {w['dns']} took "
                             "DNS accessories - the lowest penetration.")
        rec = reconcile_rows(link)
        for status, text in (
                ("In the list, no invoice found",
                 "car(s) show DNS accessories in the delivery list but no invoice was found"),
                ("Invoiced, list shows no DNS value",
                 "car(s) were invoiced but the delivery list shows no DNS value"),
                ("Amount differs", "car(s): the list's DNS value differs from the invoice")):
            n = sum(r["status"] == status for r in rec)
            if n:
                attention.append(f"{n} {text} (sheet 15).")
        follow = sum(1 for c in link.cars if not link.took_dns(c) and any(
            w in c.remarks.upper() for w in ("LATER", "WAITING", "UPDATE", "WILL")))
        if follow:
            attention.append(f"{follow} delivered car(s) without DNS accessories are marked "
                             "'later' / 'waiting' - follow up (sheet 14).")
        blank = sum(1 for c in link.cars if c.oe_list and not c.oe_value)
        if blank:
            attention.append(f"{blank} car(s) have OE accessories listed without a value, so "
                             "OE value is understated.")
        nowhere = sum(1 for c in link.cars if c.location == "(not given)")
        if nowhere:
            attention.append(f"{nowhere} car(s) in the delivery list have no location.")

    def top(title_text: str, groups: dict[str, list[float]], n: int = 5) -> None:
        # Workbook: ranked by gross profit. PDF (v0.11.1, the client's
        # request): ranked by profit margin % - the top n ARE the highest
        # margins (ties: the larger gross profit first), and the heading
        # says so. Items with no sales are left out of that ranking.
        if highest_first:
            title_text = title_text.replace("by gross profit", "by profit margin %")
            margin = lambda g: g[1][1] / g[1][0] if g[1][0] else 0
            ranked = sorted((g for g in groups.items() if g[1][0] > 0),
                            key=lambda g: (-margin(g), -g[1][1]))[:n]
        else:
            ranked = sorted(groups.items(), key=lambda g: -g[1][1])[:n]
        head(title_text, "Sales", "Gross profit", "Margin %")
        for name, (sales, gp) in ranked:
            put(name, round(sales, 2), round(gp, 2), _ratio(gp, sales),
                fmt=(MONEY, MONEY, PCT))

    products: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for l in d.lines:
        products[l.product][0] += l.sales
        products[l.product][1] += l.gross_profit
    people: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    branches: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for i in d.invoices:
        for group, key in ((people, exec_label(i.executive, i.branch)),
                           (branches, i.branch or "(no branch)")):
            group[key][0] += i.sales
            group[key][1] += i.gross_profit
    top("Top 5 products by gross profit", products)
    top("Top 5 sales executives by gross profit", people)
    top("Branches by gross profit", branches, n=20)

    packs = [i for i in d.invoices if i.package]
    near = [i for i in d.invoices if i.near_package]
    head("Packages", "")
    put("Package sales", len(packs), fmt="0")
    put("Package sales value (excl. GST)", round(sum(i.package.sales for i in packs), 2))
    put("Almost a package (one item missing)", len(near), fmt="0")
    if near:
        attention.append(f"{len(near)} invoice(s) were one item short of a package "
                         "(sheet 5).")

    incentive = round(sum(i.incentive_payable for i in d.invoices), 2)
    head("Costs", "Amount", "% of sales")
    costs = [("Direct costs (product cost + labour)", now["cogs"]),
             ("Indirect costs", now["indirect"]),
             ("Spot incentive payable (tool)", incentive)]
    if d.internal_incentive:
        costs.append(("Internal team incentive", d.internal_incentive))
    if highest_first:
        costs.sort(key=lambda c: -c[1])
    for label, amount in costs:
        put(label, amount, _ratio(amount, d.sales), fmt=(MONEY, PCT))

    if d.left_out:
        attention.append(f"{len(d.left_out)} invoice(s) are not in these figures because of "
                         "open issues on Scan review (sheet 'Not included').")
    if not d.has_inputs:
        attention.append("No indirect costs were entered on Monthly inputs: net profit is "
                         "overstated.")
    if now["net"] < 0:
        attention.append("The month shows a net loss.")
    if link is None:
        attention.append("The delivery (RTO) list was not given, so new-car penetration is "
                         "not shown.")

    if show_attention:      # left out of the PDF version (Brinda, 04-10-2026)
        head("Points needing attention", "")
        for text in attention or ["Nothing to flag this month."]:
            cell = ws.cell(state["row"], 1, "• " + text)
            cell.font = _font()
            if attention:
                cell.fill = PatternFill("solid", fgColor=AMBER_TINT)
            state["row"] += 1

    ws.column_dimensions["A"].width = 52
    for col in "BCD":
        ws.column_dimensions[col].width = 18
    _finish(ws)


def write_rto_sheets(wb, d: MonthData, link: Linked) -> list[str]:
    """Add sheets 13-17 to the workbook; returns their names."""
    writers = (penetration_sheet, missed_sheet, reconcile_sheet, source_sheet,
               scorecard_sheet)
    for name, writer in zip(RTO_SHEETS, writers):
        writer(wb.create_sheet(name), d, link)
    return list(RTO_SHEETS)
