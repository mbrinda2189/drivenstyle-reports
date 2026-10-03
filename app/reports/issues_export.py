"""
issues_export.py - Scan review issues as an Excel file (v0.6.2)
===============================================================

WHAT THIS MODULE DOES
---------------------
`export_issues(invoices, year, month, path)` writes the month's OPEN Scan
review issues to an .xlsx file that Brinda can check herself or send to
the client as a question list. It only reads the database.

    Sheet "Issues" - one row per row on the Scan review screen
        Type | As printed on the invoice | Issue | Invoices | Invoice numbers |
        What we need | Client's reply (blank, for the client to fill in)

    Sheet "Invoices affected" - one row per invoice with an open issue
        Invoice no | Date | Customer | Salesperson | Vehicle | Total (₹) |
        Issues

Why both: the first sheet is what has to be answered (one answer can clear
many invoices - e.g. "Mano Vikram", 29 invoices); the second lets the
client look up the actual bills.

"What we need" is a plain-language request that depends on the kind of
issue and why it arose (see NEEDS). Fixed issues and skipped files are not
exported - the file is about what is still open.

Same look as the reports workbook: Arial, navy headings, filters, frozen
headings, landscape, dd-mm-yyyy dates, #,##0.00 amounts.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.data.invoices_repo import InvoicesRepo, month_label
from app.reports.workbook import Col, _finish, table, title

TYPES = {"product": "Item", "salesperson": "Salesperson", "car": "Vehicle",
         "totals": "Totals"}

# (kind, reason) -> what we ask the client for. reason "" = any.
NEEDS = {
    ("product", ""): "Add this item to the Product master (selling price, cost, "
                     "labour charge), or tell us which existing item it is.",
    ("salesperson", "not_found"): "Not in the executive list: give the contact no and "
                                  "branch so they can be added, or tell us which "
                                  "existing executive this is.",
    ("salesperson", "branch"): "The executive list has this name at another branch: "
                               "confirm the correct branch.",
    ("salesperson", "several"): "More than one executive has this name: tell us who "
                                "made this sale.",
    ("salesperson", "missing"): "No salesperson on the invoice: tell us who made the sale.",
    ("car", "not_found"): "Not in the car list: give make, model and segment, or "
                          "tell us which listed car this is.",
    ("car", "several"): "Matches more than one car: tell us which one.",
    ("car", "missing"): "No vehicle on the invoice: tell us the car - or, if it was "
                        "a counter sale, mark its items 'Vehicle needed = No'.",
    ("totals", ""): "The invoice figures do not add up: please check this invoice "
                    "in Zoho.",
}


def what_we_need(kind: str, reason: str) -> str:
    return NEEDS.get((kind, reason)) or NEEDS.get((kind, ""), "")


def default_file_name(year: int, month: int) -> str:
    """DriveNStyle_Sep-2026_Issues.xlsx"""
    return f"DriveNStyle_{date(year, month, 1):%b}-{year}_Issues.xlsx"


def export_issues(invoices: InvoicesRepo, year: int, month: int,
                  path: str | Path) -> int:
    """Write the open issues to `path`. Returns the number of issues written."""
    issues = [i for i in invoices.issues(year, month) if i.status == "open"]
    order = {"product": 0, "salesperson": 1, "car": 2, "labour": 3, "totals": 4}
    issues.sort(key=lambda i: (order.get(i.kind, 9), -len(i.invoices), i.printed.lower()))
    by_no = {inv["invoice_no"]: inv for inv in invoices.invoices(year, month)}
    label = month_label(year, month)
    affected: dict[str, list[str]] = {}
    for i in issues:
        for no in i.invoices:
            affected.setdefault(no, []).append(i.message)

    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    row = title(ws, "Drive N Style – Scan review issues", label, [
        f"Exported {datetime.now():%d-%m-%Y %H:%M}. {len(issues)} open issue"
        f"{'s' if len(issues) != 1 else ''} on {len(affected)} invoice"
        f"{'s' if len(affected) != 1 else ''}. These invoices are left out of the "
        "reports until the issue is answered.",
        "Please fill in “Client's reply”. One reply clears every invoice on its row."])
    t = table(ws, row, [
        Col("Type", "type", width=13),
        Col("As printed on the invoice", "printed", width=34),
        Col("Issue", "message", width=48),
        Col("Invoices", "count", "qty", 10),
        Col("Invoice numbers", "numbers", width=34),
        Col("What we need", "need", width=48),
        Col("Client's reply", "reply", width=40),
    ], [dict(type=TYPES.get(i.kind, i.kind), printed=i.printed or "(not printed)",
             message=i.message, count=len(i.invoices), numbers=", ".join(i.invoices),
             need=what_we_need(i.kind, i.reason), reply="")
        for i in issues], filters=True, empty_text="No open issues this month.")
    _wrap(ws, t)
    _finish(ws, f"A{t['head'] + 1}")

    ws2 = wb.create_sheet("Invoices affected")
    row = title(ws2, "Invoices affected", label,
                ["Every invoice with an open issue, as read from Zoho's export."])
    rows = []
    for no in sorted(affected, key=lambda n: (by_no.get(n, {}).get("invoice_date", ""), n)):
        inv = by_no.get(no, {})
        rows.append(dict(
            no=no,
            date=date.fromisoformat(inv["invoice_date"]) if inv else None,
            customer=inv.get("customer", ""), salesperson=inv.get("salesperson", ""),
            vehicle=inv.get("vehicle", ""), total=inv.get("total", 0.0),
            issues="\n".join(affected[no])))
    t2 = table(ws2, row, [
        Col("Invoice no", "no", width=18), Col("Date", "date", "date", 12),
        Col("Customer", "customer", width=30), Col("Salesperson", "salesperson", width=24),
        Col("Vehicle", "vehicle", width=16), Col("Total (₹)", "total", "money", 14),
        Col("Issues", "issues", width=70),
    ], rows, filters=True, empty_text="No open issues this month.")
    _wrap(ws2, t2)
    _finish(ws2, f"A{t2['head'] + 1}")
    wb.save(path)
    return len(issues)


def matches_file_name() -> str:
    return f"DriveNStyle_Saved_matches_{datetime.now():%d-%m-%Y_%H%M}.xlsx"


def export_matches(invoices: InvoicesRepo, year: int, month: int,
                   path: str | Path) -> int:
    """
    v0.8.3: the "Saved matches" tab as an Excel file - every choice saved on
    Scan review (item / salesperson / vehicle matched to a master row or to
    Others, and totals differences accepted), for checking. Returns the
    number of rows written.
    """
    matches = invoices.saved_matches(year, month)
    wb = Workbook()
    ws = wb.active
    ws.title = "Saved matches"
    row = title(ws, "Scan review - saved matches", month_label(year, month), [
        "Choices saved on Scan review. They are remembered for later months. "
        "Remove a wrong one on Scan review > Saved matches.",
        f"As at {datetime.now():%d-%m-%Y %H:%M}."])
    table(ws, row, [
        Col("Type", "type", width=14),
        Col("As on the invoice", "printed", width=44),
        Col("Matched to", "target", width=44),
        Col("Applies to", "scope", width=30),
        Col("Invoices this month", "invoices", "qty", 12),
        Col("Saved on", "saved", width=18),
    ], [dict(type=m.type_label, printed=m.printed, target=m.target, scope=m.scope,
             invoices=m.invoices, saved=m.saved_on.replace("T", " "))
        for m in matches], filters=True, empty_text="No saved matches.")
    _finish(ws)
    wb.save(Path(path))
    return len(matches)


def _wrap(ws, t: dict) -> None:
    """Long texts wrap inside their cell; everything aligned to the top."""
    for r in range(t["first"], t["last"] + 1):
        for cell in ws[r]:
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal=cell.alignment.horizontal)
