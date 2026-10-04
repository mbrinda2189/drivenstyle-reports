"""
workbook.py - Writes the monthly reports workbook
=================================================

WHAT THIS MODULE DOES
---------------------
`write_workbook(data, trend, payments, reports, path)` writes one .xlsx:

    Cover                      month, when and by whom generated, key
                               figures, contents, notes
    1 Invoice profitability    one row per invoice
    2 Service vs product       Product / Service summary +
                               item-wise detail
    3 Labour                   labour by product + every line with labour
    4 Trend                    one column per scanned month (up to 12)
    5 Packages                 package sales by package + detail
    6 Vehicle-wise             by car model and by segment, average per car
    7 Spot incentive           by executive + every line (client's rule)
    8 Executive-wise sales     by executive (name and branch)
    9 High-profit products     products at or above the threshold margin
    10 Indirect vs direct      direct and indirect costs as % of sales
    11 Payment modes           received by mode, per invoice, by account
    12 Profit & loss           sales to net profit, with % of sales
    Not included               invoices left out (open issues) and
                               skipped files, with reasons

Only the reports ticked on the Generate page are written; Cover and Not
included are always there.

LIVE FORMULAS
-------------
Figures that come from the invoices and masters (sales, costs, GST ...)
are written as values. Everything worked out from them - gross profit,
margin %, averages, totals, shares, summaries of a detail table (SUMIF) -
is an Excel formula, so the workbook stays consistent if a figure is
corrected by hand. Excel recalculates everything when the file is opened.

LOOK
----
Arial throughout; headings in the tool's navy; totals bold with a rule
above; amounts as 12,345.00; percentages as 12.3%; dates dd-mm-yyyy;
headings frozen; filters on detail tables; landscape, fit to page width
when printed.
"""

from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.data.payments_io import Payment
from app.reports.data import MonthData

FONT = "Arial"
NAVY, SLATE, AMBER, AMBER_TINT = "0B2545", "5B6B82", "B7791F", "FBF1DE"
# QTY is "General": "#,##0.##" showed whole quantities with a trailing dot
# ("9.", "108.") - fixed in v0.6.2. General shows 9 and 1.5 as they are.
MONEY, PCT, DATE, QTY = "#,##0.00", "0.0%", "DD-MM-YYYY", "General"
thin = Side(style="thin", color="DCE4EF")
rule = Side(style="thin", color=NAVY)

# Sheet names, in the order of sample_data.REPORTS.
SHEETS = OrderedDict([
    ("Invoice-wise profitability", "1 Invoice profitability"),
    ("Service vs product profitability", "2 Service vs product"),
    ("Labour calculation", "3 Labour"),
    ("Trend analysis (month on month)", "4 Trend"),
    ("Basic package analysis", "5 Packages"),
    ("Vehicle-wise average per car", "6 Vehicle-wise"),
    ("Spot incentive calculation", "7 Spot incentive"),
    ("Executive-wise sales", "8 Executive-wise sales"),
    ("High-profit product sales", "9 High-profit products"),
    ("Indirect vs direct cost %", "10 Indirect vs direct"),
    ("Payment mode analysis", "11 Payment modes"),
    ("Profit & loss", "12 Profit & loss"),
])


# ---------------------------------------------------------------------------
# Small writing helpers
# ---------------------------------------------------------------------------
@dataclass
class Col:
    """
    One table column.
    key      field of the row dict to write (None for a formula column)
    formula  e.g. "={sales}{r}-{cost}{r}" - {name} is the letter of the
             column with that key, {r} the row
    total    "sum", a formula template (with {r} = total row), or None
    """
    header: str
    key: str | None = None
    kind: str = "text"            # text / money / pct / date / qty
    width: float = 14
    formula: str | None = None
    total: str | None = None


def _font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def title(ws: Worksheet, text: str, month: str, notes: list[str] = ()) -> int:
    """Sheet title, month and notes. Returns the next free row."""
    ws["A1"] = text
    ws["A1"].font = _font(True, NAVY, 14)
    ws["A2"] = month
    ws["A2"].font = _font(color=SLATE, italic=True)
    row = 3
    for note in notes:
        ws.cell(row, 1, note).font = _font(color=SLATE, size=9)
        row += 1
    return row + 1


def banner(ws: Worksheet, row: int, text: str, span: int = 8) -> int:
    """An amber note across the sheet (e.g. PROVISIONAL)."""
    ws.cell(row, 1, text).font = _font(True, AMBER)
    for c in range(1, span + 1):
        ws.cell(row, c).fill = PatternFill("solid", fgColor=AMBER_TINT)
    return row + 2


def section(ws: Worksheet, row: int, text: str) -> int:
    ws.cell(row, 1, text).font = _font(True, NAVY, 11)
    return row + 1


def table(ws: Worksheet, row: int, cols: list[Col], rows: list[dict],
          total_label: str = "Total", filters: bool = False,
          empty_text: str = "Nothing to show for this month.") -> dict:
    """
    Write a table with a heading row starting at `row`. Returns
    {"head": row, "first": first data row, "last": last data row,
     "total": total row or None, "next": next free row,
     "letter": {key or header: column letter}}.
    """
    letter = {}
    for i, c in enumerate(cols, start=1):
        letter[c.key or c.header] = get_column_letter(i)
        cell = ws.cell(row, i, c.header)
        cell.font = _font(True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="right" if c.kind in ("money", "pct", "qty")
                                   else "left")
        current = ws.column_dimensions[get_column_letter(i)].width or 0
        ws.column_dimensions[get_column_letter(i)].width = max(current, c.width)
    ws.row_dimensions[row].height = 30
    head = row
    if not rows:
        ws.cell(row + 1, 1, empty_text).font = _font(italic=True, color=SLATE)
        return {"head": head, "first": row + 1, "last": row, "total": None,
                "next": row + 3, "letter": letter}

    for data in rows:
        row += 1
        for i, c in enumerate(cols, start=1):
            if c.formula:
                value = c.formula.format(r=row, **letter)
            else:
                value = data.get(c.key, "")
            cell = ws.cell(row, i, value)
            _style(cell, c.kind)
            cell.border = Border(bottom=thin)
    first, last = head + 1, row

    total_row = None
    if any(c.total for c in cols):
        row += 1
        total_row = row
        ws.cell(row, 1, total_label)
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row, i)
            if c.total == "sum":
                col = letter[c.key or c.header]
                cell.value = f"=SUM({col}{first}:{col}{last})"
            elif c.total:
                cell.value = c.total.format(r=row, first=first, last=last, **letter)
            _style(cell, c.kind)
            cell.font = _font(True)
            cell.border = Border(top=rule)
    if filters:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{last}"
    return {"head": head, "first": first, "last": last, "total": total_row,
            "next": row + 3, "letter": letter}


def _style(cell, kind: str) -> None:
    cell.font = _font()
    if kind == "money":
        cell.number_format = MONEY
    elif kind == "pct":
        cell.number_format = PCT
    elif kind == "date":
        cell.number_format = DATE
        cell.alignment = Alignment(horizontal="left")
    elif kind == "qty":
        cell.number_format = QTY


def _finish(ws: Worksheet, freeze: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # v0.9.1 - printing and the PDF (pdf_export.py): A4, narrow margins,
    # "sheet name - page x of y" at the foot of every page, and on long
    # sheets (those with frozen headings) the heading row repeated at the
    # top of each further page.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "&A  -  page &P of &N"
    ws.oddFooter.center.size = 8
    if freeze:
        head = int("".join(ch for ch in freeze if ch.isdigit())) - 1
        if head >= 1:
            ws.print_title_rows = f"{head}:{head}"


def exec_label(name: str, branch: str) -> str:
    """Executives are shown with their branch: two PRAVEENs are different people."""
    return f"{name} – {branch}" if branch else name


GP = "={sales}{r}-{cost}{r}-{labour}{r}"
MARGIN = '=IF({sales}{r}=0,"",{gp}{r}/{sales}{r})'
MARGIN_TOTAL = '=IF({sales}{r}=0,"",{gp}{r}/{sales}{r})'


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------
def write_workbook(data: MonthData, trend: list[MonthData],
                   payments: list[Payment] | None, reports: list[str],
                   path: str | Path, user: str = "", rto: list | None = None,
                   previous: MonthData | None = None,
                   rto_by_month: dict | None = None) -> Path:
    """
    Write the workbook for `data` (see module notes). Returns the path.
    v0.9.0: `rto` = the dealership's delivery list (app/data/rto_list.py);
    when given, sheets 13-17 are added. A "Summary" sheet (executive
    summary) is always written, comparing with `previous` (the month
    before) when there is one. See app/reports/rto_reports.py.
    """
    from app.reports import rto_reports as rr      # (it imports this module)
    from app.reports.trend import trend_sheet      # (so does this)
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    written = []
    writers = {
        "Invoice-wise profitability": lambda ws: invoice_sheet(ws, data),
        "Service vs product profitability": lambda ws: category_sheet(ws, data),
        "Labour calculation": lambda ws: labour_sheet(ws, data),
        "Trend analysis (month on month)": lambda ws: trend_sheet(ws, data, trend, rto_by_month),
        "Basic package analysis": lambda ws: package_sheet(ws, data),
        "Vehicle-wise average per car": lambda ws: vehicle_sheet(ws, data),
        "Spot incentive calculation": lambda ws: incentive_sheet(ws, data),
        "Executive-wise sales": lambda ws: executive_sheet(
            ws, data, SHEETS["Spot incentive calculation"]
            if "Spot incentive calculation" in reports else None),
        "High-profit product sales": lambda ws: high_profit_sheet(ws, data),
        "Indirect vs direct cost %": lambda ws: cost_split_sheet(ws, data),
        "Payment mode analysis": lambda ws: payment_sheet(ws, data, payments),
        "Profit & loss": lambda ws: pnl_sheet(ws, data),
    }
    for name, sheet_name in SHEETS.items():
        if name in reports:
            ws = wb.create_sheet(sheet_name)
            writers[name](ws)
            written.append(sheet_name)
    link = rr.Linked(data, rto) if rto is not None else None
    if link is not None:
        written += rr.write_rto_sheets(wb, data, link)
    left = wb.create_sheet("Not included")
    not_included_sheet(left, data)
    summary = wb.create_sheet(rr.SUMMARY_SHEET, 1)          # right after the Cover
    rr.summary_sheet(summary, data, previous, link, user)
    cover_sheet(cover, data, [rr.SUMMARY_SHEET] + written, user)
    wb.calculation.fullCalcOnLoad = True       # Excel works out every formula
    path = Path(path)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------
def cover_sheet(ws: Worksheet, d: MonthData, sheets: list[str], user: str,
                pdf: bool = False) -> None:
    """`pdf` (v0.10.0): the cover of the PDF version lists only its own pages."""
    ws["A1"] = "Drive N Style – Monthly reports"
    ws["A1"].font = _font(True, NAVY, 16)
    ws["A2"] = d.label
    ws["A2"].font = _font(True, SLATE, 12)
    ws["A3"] = (f"Generated on {datetime.now():%d-%m-%Y %H:%M}"
                + (f" by {user}" if user else ""))
    ws["A3"].font = _font(color=SLATE, size=9)

    left_total = sum(x.total for x in d.left_out)
    facts = [
        ("Invoices included", len(d.invoices), None),
        ("Sales (excluding GST)", d.sales, MONEY),
        ("Gross profit", d.gross_profit, MONEY),
        ("Invoices not included (open issues)", len(d.left_out), None),
        ("Amount of invoices not included", left_total, MONEY),
        ("Files skipped", len(d.skipped_files), None),
        ("Rs. 1 labour marker lines ignored", d.marker_lines, None),
        ("Value of those lines (not in sales)", d.marker_value, MONEY),
    ]
    row = 5
    for label, value, fmt in facts:
        ws.cell(row, 1, label).font = _font()
        c = ws.cell(row, 2, value)
        c.font = _font(True)
        if fmt:
            c.number_format = fmt
        row += 1

    names = sheets + ([] if pdf else ["Not included"])
    if names:                      # (the PDF version is one sheet: no contents)
        row = section(ws, row + 1, "Contents")
    for s in names:
        c = ws.cell(row, 1, s)
        c.hyperlink = f"#'{s}'!A1"
        c.font = Font(name=FONT, color="1F5FBF", underline="single")
        row += 1

    row = section(ws, row + 1, "Notes")
    notes = [
        "Sales are after discount and exclude GST. Invoices that show no GST count in full as sales.",
        "Product cost and labour come from the Product master, at the rates that applied on each invoice date.",
        "Labour is taken from the product (Labour involved + labour charge). The Rs. 1 'Labour Charges …' / "
        "'Labour - …' lines only show that labour was done: they are ignored, and their small value is "
        "not in sales, so sales are slightly below the invoice totals.",
        "Invoices with open issues on Scan review are left out of every report and listed on 'Not included'.",
    ]
    if not d.has_inputs:
        notes.append("No indirect costs were entered for this month (Monthly inputs): they are shown as zero.")
    for n in notes:
        ws.cell(row, 1, "• " + n).font = _font(size=9)
        row += 1
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    _finish(ws)


# ---------------------------------------------------------------------------
# 1 Invoice-wise profitability
# ---------------------------------------------------------------------------
def invoice_sheet(ws: Worksheet, d: MonthData) -> None:
    row = title(ws, "Invoice-wise profitability", d.label,
                ["Sales are after discount, before GST. Gross profit = sales − product cost − labour."])
    cols = [
        Col("Invoice no", "invoice_no", width=17),
        Col("Date", "date", "date", 11),
        Col("Customer", "customer", width=24),
        Col("Customer type", "customer_type", width=12),
        Col("Salesperson", "executive", width=20),
        Col("Car", "car", width=16),
        Col("Items", "items", "qty", 7, total="sum"),
        Col("Sales", "sales", "money", 13, total="sum"),
        Col("Discount", "discount", "money", 11, total="sum"),
        Col("GST", "gst", "money", 11, total="sum"),
        Col("Rounding", "rounding", "money", 9, total="sum"),
        Col("Invoice total", "total", "money", 13, total="sum"),
        Col("Product cost", "cost", "money", 13, total="sum"),
        Col("Labour", "labour", "money", 11, total="sum"),
        Col("Gross profit", "gp", "money", 13, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 9, formula=MARGIN, total=MARGIN_TOTAL),
    ]
    rows = [dict(invoice_no=i.invoice_no, date=i.invoice_date, customer=i.customer,
                 customer_type=i.customer_type,
                 # "Others" (v0.7.1) also shows the name printed on the invoice
                 executive=exec_label(i.executive, i.branch)
                 + (f" ({i.printed_executive})" if i.printed_executive else ""),
                 car=i.car + (f" ({i.printed_car})" if i.printed_car else ""),
                 items=i.items, sales=i.sales, discount=i.discount, gst=i.gst,
                 rounding=i.rounding, total=i.total, cost=i.cost, labour=i.labour)
            for i in sorted(d.invoices, key=lambda x: (x.invoice_date, x.invoice_no))]
    t = table(ws, row, cols, rows, filters=True)
    _finish(ws, f"B{t['head'] + 1}")


# ---------------------------------------------------------------------------
# 2 Service vs product
# ---------------------------------------------------------------------------
def _by_product(lines) -> list[dict]:
    """Lines added up per product (and category)."""
    out: dict[tuple, dict] = {}
    for l in lines:
        k = (l.category, l.product)
        r = out.setdefault(k, dict(category=l.category, product=l.product, qty=0.0,
                                   sales=0.0, cost=0.0, labour=0.0))
        r["qty"] += l.qty
        r["sales"] += l.sales
        r["cost"] += l.cost
        r["labour"] += l.labour
    for r in out.values():
        for k in ("sales", "cost", "labour"):
            r[k] = round(r[k], 2)
    return list(out.values())


def category_sheet(ws: Worksheet, d: MonthData) -> None:
    row = title(ws, "Service vs product profitability", d.label,
                ["Category comes from the Product master. The Rs. 1 labour marker lines "
                 "are ignored."])
    categories = ["Product", "Service"]
    detail = sorted(_by_product(d.lines),
                    key=lambda r: (categories.index(r["category"])
                                   if r["category"] in categories else 9, -r["sales"]))
    # Summary first; it adds up the detail table below with SUMIF.
    detail_head = row + len(categories) + 5
    first, last = detail_head + 1, detail_head + max(1, len(detail))

    def sumif(col):
        return f"=SUMIF($A${first}:$A${last},$A{{r}},{col}${first}:{col}${last})"
    row = section(ws, row, "Summary")
    summary_cols = [
        Col("Category", "category", width=16),
        Col("Qty", "qty", "qty", 9, formula=sumif("$C"), total="sum"),
        Col("Sales", "sales", "money", 14, formula=sumif("$D"), total="sum"),
        Col("Product cost", "cost", "money", 14, formula=sumif("$E"), total="sum"),
        Col("Labour", "labour", "money", 12, formula=sumif("$F"), total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 10, formula=MARGIN, total=MARGIN_TOTAL),
        Col("Share of sales", "share", "pct", 10,
            formula='=IF(SUM({sales}$%d:{sales}$%d)=0,"",{sales}{r}/SUM({sales}$%d:{sales}$%d))'
            % (row + 1, row + len(categories), row + 1, row + len(categories)),
            total='=IF({sales}{r}=0,"",1)'),
    ]
    table(ws, row, summary_cols, [dict(category=c) for c in categories])

    row = section(ws, detail_head - 1, "Item-wise detail")
    detail_cols = [
        Col("Category", "category", width=16),
        Col("Product / service", "product", width=40),
        Col("Qty", "qty", "qty", 9, total="sum"),
        Col("Sales", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 10, formula=MARGIN, total=MARGIN_TOTAL),
    ]
    table(ws, row, detail_cols, detail, filters=True)
    _finish(ws)


# ---------------------------------------------------------------------------
# 3 Labour
# ---------------------------------------------------------------------------
def labour_sheet(ws: Worksheet, d: MonthData) -> None:
    lines = [l for l in d.lines if l.labour > 0]
    row = title(ws, "Labour calculation", d.label,
                ["Labour cost = labour charge in the Product master × quantity, for products "
                 "marked 'Labour involved'.",
                 f"The Rs. 1 labour marker lines on the invoices ({d.marker_lines} this month) "
                 "are not used: labour comes from the product."])
    products = sorted({l.product for l in lines})
    detail_head = row + len(products) + 5
    first, last = detail_head + 1, detail_head + max(1, len(lines))

    row = section(ws, row, "By product")
    sumif = lambda col: f"=SUMIF($D${first}:$D${last},$A{{r}},{col}${first}:{col}${last})"
    table(ws, row, [
        Col("Product / service", "product", width=40),
        Col("Qty", "qty", "qty", 9, formula=sumif("$E"), total="sum"),
        Col("Labour cost", "labour_total", "money", 14, formula=sumif("$G"), total="sum"),
    ], [dict(product=p) for p in products])

    row = section(ws, detail_head - 1, "Every line with labour")
    table(ws, row, [
        Col("Invoice no", "invoice_no", width=17),
        Col("Date", "date", "date", 11),
        Col("Salesperson", "executive", width=22),
        Col("Product / service", "product", width=40),
        Col("Qty", "qty", "qty", 9, total="sum"),
        Col("Labour charge / unit", "rate", "money", 14),
        Col("Labour cost", "labour_cost", "money", 14, formula="={qty}{r}*{rate}{r}", total="sum"),
    ], [dict(invoice_no=l.invoice_no, date=l.invoice_date,
             executive=exec_label(l.executive, l.branch), product=l.product,
             qty=l.qty, rate=l.labour_rate) for l in lines], filters=True)
    _finish(ws)


# ---------------------------------------------------------------------------
# 4 Trend
# ---------------------------------------------------------------------------
# (rebuilt in v0.10.0 - see app/reports/trend.py: more rows, change from the
# previous month, sales by branch, top products and charts)


# ---------------------------------------------------------------------------
# 5 Packages
# ---------------------------------------------------------------------------
def package_sheet(ws: Worksheet, d: MonthData) -> None:
    """
    v0.8.0: a package sale = an invoice carrying EVERY item of a package
    (Packages master; see app/reports/packages.py). Values:
      list value    the package items at Zoho's selling price
      coupon value  the package's final value (Incentive master, Bill value)
      amount billed what the customer paid for the package items, with GST
    """
    sold = [i for i in d.invoices if i.package]
    near = [i for i in d.invoices if i.near_package]
    row = title(ws, "Basic package analysis", d.label, [
        "A package sale is an invoice that has every item of a package (Packages master). "
        "Other items on the same invoice are normal sales and are not counted here.",
        "List value = the package items at the Product master's selling price. Coupon value = "
        "the package's final value (Bill value in the Incentive master). Amount billed "
        "includes GST.",
        f"Result as at {datetime.now():%d-%m-%Y %H:%M}."])

    groups: dict[str, dict] = {}
    for i in sold:
        k = i.package
        g = groups.setdefault(k.package, dict(package=k.package, qty=0, list=0.0,
                                              coupon=0.0, billed=0.0, sales=0.0,
                                              cost=0.0, labour=0.0))
        g["qty"] += 1
        for key, value in (("list", k.list_value), ("coupon", k.coupon_value),
                           ("billed", k.billed), ("sales", k.sales), ("cost", k.cost),
                           ("labour", k.labour)):
            g[key] = round(g[key] + value, 2)
    row = section(ws, row, "By package")
    t = table(ws, row, [
        Col("Package", "package", width=30),
        Col("Times sold", "qty", "qty", 10, total="sum"),
        Col("List value", "list", "money", 14, total="sum"),
        Col("Coupon value", "coupon", "money", 14, total="sum"),
        Col("Amount billed", "billed", "money", 14, total="sum"),
        Col("Sales (excl. GST)", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 10, formula=MARGIN, total=MARGIN_TOTAL),
    ], sorted(groups.values(), key=lambda g: -g["sales"]),
        empty_text="No package sales this month.")

    row = section(ws, t["next"], "Package invoices")
    payable = ('=IF({coupon}{r}=0,0,{inc}{r}*MIN(1,{billed}{r}/{coupon}{r}))')
    t = table(ws, row, [
        Col("Invoice no", "invoice_no", width=17),
        Col("Date", "date", "date", 11),
        Col("Salesperson", "executive", width=24),
        Col("Car", "car", width=18),
        Col("Package", "package", width=30),
        Col("Items", "items", "qty", 7),
        Col("List value", "list", "money", 13, total="sum"),
        Col("Coupon value", "coupon", "money", 13, total="sum"),
        Col("Amount billed", "billed", "money", 13, total="sum"),
        Col("Discount given", "given", "money", 13,
            formula="={list}{r}-{billed}{r}", total="sum"),
        Col("Coupon discount", "allowed", "money", 13,
            formula="={list}{r}-{coupon}{r}", total="sum"),
        Col("Sales (excl. GST)", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 13, total="sum"),
        Col("Labour", "labour", "money", 11, total="sum"),
        Col("Gross profit", "gp", "money", 13, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 9, formula=MARGIN, total=MARGIN_TOTAL),
        Col("Package incentive", "inc", "money", 13, total="sum"),
        Col("Incentive payable", "payable", "money", 14, formula=payable, total="sum"),
    ], [dict(invoice_no=i.invoice_no, date=i.invoice_date,
             executive=exec_label(i.executive, i.branch), car=i.car,
             package=i.package.package, items=len(i.package.lines),
             list=i.package.list_value, coupon=i.package.coupon_value,
             billed=i.package.billed, sales=i.package.sales, cost=i.package.cost,
             labour=i.package.labour, inc=i.package.incentive) for i in sold],
        filters=True, empty_text="No package sales this month.")

    row = section(ws, t["next"], "Package items as billed")
    t = table(ws, row, [
        Col("Invoice no", "invoice_no", width=17),
        Col("Package", "package", width=30),
        Col("Product as billed", "product", width=40),
        Col("Qty", "qty", "qty", 8),
        Col("List price", "list", "money", 13),
        Col("Sales (excl. GST)", "sales", "money", 14, total="sum"),
    ], [dict(invoice_no=i.invoice_no, package=i.package.package, product=l.product,
             qty=l.qty, list=l.list_price, sales=l.sales)
        for i in sold for l in i.package.lines],
        empty_text="No package sales this month.")

    row = section(ws, t["next"], "Almost a package (one item missing) - check with the client")
    table(ws, row, [
        Col("Invoice no", "invoice_no", width=17),
        Col("Date", "date", "date", 11),
        Col("Salesperson", "executive", width=24),
        Col("Car", "car", width=18),
        Col("Package", "package", width=30),
        Col("Item missing", "missing", width=30),
        Col("Invoice total", "total", "money", 13),
    ], [dict(invoice_no=i.invoice_no, date=i.invoice_date,
             executive=exec_label(i.executive, i.branch), car=i.car,
             package=i.near_package, missing=i.near_missing, total=i.total)
        for i in near], empty_text="None this month.")
    _finish(ws)


# ---------------------------------------------------------------------------
# 6 Vehicle-wise
# ---------------------------------------------------------------------------
def vehicle_sheet(ws: Worksheet, d: MonthData) -> None:
    row = title(ws, "Vehicle-wise average per car", d.label,
                ["Each invoice is one car. Car and segment come from the Car master.",
                 "Counter sales (items marked 'Vehicle needed = No', no car on the invoice) "
                 "are shown as one row, 'Counter sale (no vehicle)'."])
    cars: dict[str, dict] = {}
    for i in d.invoices:
        c = cars.setdefault(i.car, dict(car=i.car, segment=i.segment, cars=0,
                                        sales=0.0, cost=0.0, labour=0.0))
        c["cars"] += 1
        c["sales"] += i.sales
        c["cost"] += i.cost
        c["labour"] += i.labour
    car_rows = sorted(cars.values(), key=lambda c: -c["sales"])
    for c in car_rows:
        for k in ("sales", "cost", "labour"):
            c[k] = round(c[k], 2)
    segments = sorted({c["segment"] or "(no segment)" for c in car_rows})
    for c in car_rows:
        c["segment"] = c["segment"] or "(no segment)"

    avg = '=IF({cars}{r}=0,"",{sales}{r}/{cars}{r})'
    avg_gp = '=IF({cars}{r}=0,"",{gp}{r}/{cars}{r})'
    seg_head = row + 1
    car_head = seg_head + len(segments) + 5
    first, last = car_head + 1, car_head + max(1, len(car_rows))
    sumif = lambda col: f"=SUMIF($B${first}:$B${last},$A{{r}},{col}${first}:{col}${last})"

    row = section(ws, row, "By segment")
    table(ws, row, [
        Col("Segment", "segment", width=22),
        Col("Cars", "cars", "qty", 8, formula=sumif("$C"), total="sum"),
        Col("Sales", "sales", "money", 14, formula=sumif("$D"), total="sum"),
        Col("Product cost", "cost", "money", 14, formula=sumif("$E"), total="sum"),
        Col("Labour", "labour", "money", 12, formula=sumif("$F"), total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Average sales per car", "avg", "money", 14, formula=avg, total=avg),
        Col("Average profit per car", "avg_gp", "money", 14, formula=avg_gp, total=avg_gp),
    ], [dict(segment=s) for s in segments])

    row = section(ws, car_head - 1, "By car model")
    table(ws, row, [
        Col("Car", "car", width=22),
        Col("Segment", "segment", width=16),
        Col("Cars", "cars", "qty", 8, total="sum"),
        Col("Sales", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Average sales per car", "avg", "money", 14, formula=avg, total=avg),
        Col("Average profit per car", "avg_gp", "money", 14, formula=avg_gp, total=avg_gp),
    ], car_rows, filters=True)
    _finish(ws)


# ---------------------------------------------------------------------------
# 7 Spot incentive (rule confirmed by the client, 30-09-2026)
# ---------------------------------------------------------------------------
INTERNAL_TEAM = "Internal team"
NO_INCENTIVE_NOTE = ("No incentive for executives marked 'Gets incentive = No' in the Sales "
                     "executive master: ")


def incentive_sheet(ws: Worksheet, d: MonthData) -> None:
    # v0.8.0: the items of a package carry no incentive of their own; the
    # invoice gets one row with the package incentive instead. Its bill
    # value is the coupon's final value and the amount billed is what the
    # customer paid for the package items - the same confirmed rule applies.
    lines = [dict(executive=exec_label(l.executive, l.branch), invoice_no=l.invoice_no,
                  date=l.invoice_date, product=l.product, group=l.incentive_group,
                  qty=l.qty, inc=l.incentive_amount, bill=l.bill_value, billed=l.billed)
             for l in d.lines if l.incentive_group]
    lines += [dict(executive=exec_label(i.executive, i.branch), invoice_no=i.invoice_no,
                   date=i.invoice_date,
                   product=f"Package ({len(i.package.lines)} items)",
                   group=i.package.package, qty=1, inc=i.package.incentive,
                   bill=i.package.coupon_value, billed=i.package.billed)
              for i in d.invoices if i.package and i.incentive_allowed]
    # v0.12.0 - internal team incentive: one row per qualifying line under
    # the name "Internal team" (one line in "By executive", not split by
    # person). It is paid in full, so its bill value is set to what was
    # billed and the rule's ratio is exactly 1.
    lines += [dict(executive=INTERNAL_TEAM, invoice_no=l.invoice_no, date=l.invoice_date,
                   product=l.product, group="Internal team incentive", qty=l.qty,
                   inc=l.internal_incentive, bill=round(l.billed / l.qty, 2) if l.qty else 0,
                   billed=l.billed)
              for l in d.lines if l.internal_incentive and l.qty]
    lines.sort(key=lambda r: (r["executive"] == INTERNAL_TEAM, r["executive"], r["date"],
                              r["invoice_no"]))
    row = title(ws, "Spot incentive calculation", d.label, [
        "Rule used: payable = incentive × qty × (amount billed ÷ (bill value × qty)), "
        "never more than the full incentive.",
        "Amount billed = the line after its share of the discount, including GST. "
        "Incentive and bill value: Incentive master, on the invoice date.",
        "Package sales: one row per invoice with the package incentive; the package's items "
        "earn no separate incentive.",
        "'By executive': each executive's incentive payable is rounded up to the next Rs. 10. "
        "The lines below show the exact amounts."]
        + ([NO_INCENTIVE_NOTE + ", ".join(d.no_incentive) + "."] if d.no_incentive else [])
        + (["'Internal team': the product's Internal incentive (Product master) × quantity, "
            "paid in full in addition to the salesperson's incentive."]
           if d.internal_incentive else []))
    execs = sorted({r["executive"] for r in lines}, key=lambda e: (e == INTERNAL_TEAM, e))
    detail_head = row + len(execs) + 5
    first, last = detail_head + 1, detail_head + max(1, len(lines))
    sumif = lambda col: f"=SUMIF($A${first}:$A${last},$A{{r}},{col}${first}:{col}${last})"

    row = section(ws, row, "By executive")
    table(ws, row, [
        Col("Executive", "executive", width=26),
        Col("Lines", "n", "qty", 8,
            formula=f"=COUNTIF($A${first}:$A${last},$A{{r}})", total="sum"),
        Col("Full incentive", "full", "money", 14, formula=sumif("$J"), total="sum"),
        # v0.12.1: each executive's total is rounded UP to the next Rs. 10
        # (1,492 -> 1,500); the lines below keep the exact figures.
        Col("Incentive payable", "payable", "money", 14,
            formula="=CEILING(ROUND(" + sumif("$L")[1:] + ",2),10)", total="sum"),
        Col("Reduced by discounts", "reduced", "money", 14,
            formula="={full}{r}-{payable}{r}", total="sum"),
    ], [dict(executive=e) for e in execs])

    row = section(ws, detail_head - 1, "Every line with an incentive")
    table(ws, row, [
        Col("Executive", "executive", width=26),            # A
        Col("Invoice no", "invoice_no", width=17),          # B
        Col("Date", "date", "date", 11),                    # C
        Col("Product as billed", "product", width=34),      # D
        Col("Incentive group", "group", width=24),          # E
        Col("Qty", "qty", "qty", 7),                        # F
        Col("Incentive / unit", "inc", "money", 12),        # G
        Col("Bill value / unit", "bill", "money", 12),      # H
        Col("Amount billed", "billed", "money", 13),        # I
        Col("Full incentive", "full", "money", 13,          # J
            formula="={inc}{r}*{qty}{r}", total="sum"),
        Col("Discount %", "disc", "pct", 10,                # K
            formula='=IF({bill}{r}*{qty}{r}=0,"",MAX(0,1-{billed}{r}/({bill}{r}*{qty}{r})))'),
        Col("Incentive payable", "payable", "money", 14,    # L
            formula='=IF({bill}{r}*{qty}{r}=0,0,{full}{r}*MIN(1,{billed}{r}/({bill}{r}*{qty}{r})))',
            total="sum"),
    ], lines, filters=True, empty_text="No lines with an incentive group this month.")
    _finish(ws)


# ---------------------------------------------------------------------------
# 8 Executive-wise sales
# ---------------------------------------------------------------------------
def executive_sheet(ws: Worksheet, d: MonthData, incentive_sheet_name: str | None) -> None:
    row = title(ws, "Executive-wise sales", d.label,
                ["Executives are shown with their branch (two people may share a name)."])
    execs: dict[str, dict] = {}
    for i in d.invoices:
        key = exec_label(i.executive, i.branch)
        e = execs.setdefault(key, dict(executive=key, branch=i.branch, invoices=0, items=0.0,
                                       sales=0.0, cost=0.0, labour=0.0))
        e["invoices"] += 1
        e["items"] += i.items
        e["sales"] += i.sales
        e["cost"] += i.cost
        e["labour"] += i.labour
    rows = sorted(execs.values(), key=lambda e: -e["sales"])
    for e in rows:
        for k in ("sales", "cost", "labour"):
            e[k] = round(e[k], 2)
    cols = [
        Col("Executive", "executive", width=26),
        Col("Branch", "branch", width=12),
        Col("Invoices", "invoices", "qty", 9, total="sum"),
        Col("Items", "items", "qty", 8, total="sum"),
        Col("Sales", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 9, formula=MARGIN, total=MARGIN_TOTAL),
        Col("Average bill", "avg", "money", 13,
            formula='=IF({invoices}{r}=0,"",{sales}{r}/{invoices}{r})',
            total='=IF({invoices}{r}=0,"",{sales}{r}/{invoices}{r})'),
    ]
    if incentive_sheet_name:
        ref = f"'{incentive_sheet_name}'"
        cols.append(Col("Incentive payable", "incentive", "money", 16,
                        formula=f"=CEILING(ROUND(SUMIF({ref}!$A:$A,$A{{r}},{ref}!$L:$L),2),10)",
                        total="sum"))        # rounded up to Rs. 10, as on sheet 7
    t = table(ws, row, cols, rows, filters=True)
    _finish(ws, f"B{t['head'] + 1}")


# ---------------------------------------------------------------------------
# 9 High-profit products
# ---------------------------------------------------------------------------
def high_profit_sheet(ws: Worksheet, d: MonthData) -> None:
    th = d.threshold
    row = title(ws, "High-profit product sales", d.label,
                [f"Products with a margin of {th:g}% or more this month (threshold set on "
                 "Monthly inputs), highest gross profit first."])
    rows = []
    for r in _by_product(d.lines):
        gp = r["sales"] - r["cost"] - r["labour"]
        if r["sales"] > 0 and gp / r["sales"] * 100 >= th - 1e-9:
            r["_gp"] = gp
            rows.append(r)
    rows.sort(key=lambda r: -r["_gp"])
    table(ws, row, [
        Col("Product / service", "product", width=40),
        Col("Category", "category", width=12),
        Col("Qty", "qty", "qty", 8, total="sum"),
        Col("Sales", "sales", "money", 14, total="sum"),
        Col("Product cost", "cost", "money", 14, total="sum"),
        Col("Labour", "labour", "money", 12, total="sum"),
        Col("Gross profit", "gp", "money", 14, formula=GP, total="sum"),
        Col("Margin %", "margin", "pct", 9, formula=MARGIN, total=MARGIN_TOTAL),
    ], rows, filters=True,
        empty_text=f"No product reached a {th:g}% margin this month.")
    _finish(ws)


# ---------------------------------------------------------------------------
# 10 Indirect vs direct, 12 Profit & loss
# ---------------------------------------------------------------------------
def _statement(ws: Worksheet, row: int, d: MonthData, net: bool,
               largest_first: bool = False) -> None:
    """
    Shared layout of sheets 10 and 12: Particulars | Amount | % of sales.
    `largest_first` (v0.10.5, PDF version only): the indirect cost heads -
    typed and automatic together - are listed from the largest to the
    smallest, so "% of sales" runs downwards. The workbook keeps the order
    entered on Monthly inputs.
    """
    for c, (text, width) in enumerate((("Particulars", 38), ("Amount", 16),
                                       ("% of sales", 12)), start=1):
        cell = ws.cell(row, c, text)
        cell.font = _font(True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="right" if c > 1 else "left")
        ws.column_dimensions[get_column_letter(c)].width = width
    sales_row = row + 1
    lines: list[tuple] = []          # (label, value or formula, bold, rule, pct)

    def put(label, value, bold=False, top=False, pct=True, indent=0):
        nonlocal row
        row += 1
        a = ws.cell(row, 1, label)
        a.font = _font(bold)
        a.alignment = Alignment(indent=indent)
        b = ws.cell(row, 2, value)
        b.font, b.number_format = _font(bold), MONEY
        if pct:
            p = ws.cell(row, 3, f'=IF($B${sales_row}=0,"",B{row}/$B${sales_row})')
            p.font, p.number_format = _font(bold), PCT
        if top:
            for c in (1, 2, 3):
                ws.cell(row, c).border = Border(top=rule)
        return row

    cost = round(sum(i.cost for i in d.invoices), 2)
    labour = round(sum(i.labour for i in d.invoices), 2)
    put("Sales (excluding GST)", d.sales, bold=True)
    row += 1
    ws.cell(row, 1, "Direct costs").font = _font(True, NAVY)
    c1 = put("Product cost", cost, indent=1)
    c2 = put("Labour", labour, indent=1)
    direct = put("Total direct costs", f"=B{c1}+B{c2}", bold=True, top=True)
    gp = put("Gross profit", f"=B{sales_row}-B{direct}", bold=True) if net else None
    row += 1
    ws.cell(row, 1, "Indirect costs").font = _font(True, NAVY)
    # (label, value or formula, amount for sorting). v0.8.1: the automatic
    # heads are worked out from COGS (= total direct costs: product cost +
    # labour) - see AUTO_INDIRECT in reports/data.py.
    heads = [(h, a, a) for h, a in d.entered_indirect] or \
        [("(none entered on Monthly inputs)", 0.0, 0.0)]
    heads += [(f"{head} ({pct * 100:g}% of COGS)", f"=ROUND(B{direct}*{pct:.6g},2)", amount)
              for head, pct, amount in d.auto_indirect]
    if largest_first:
        heads.sort(key=lambda h: -h[2])
    head_rows = [put(h, value, indent=1) for h, value, _ in heads]
    indirect = put("Total indirect costs", f"=SUM(B{head_rows[0]}:B{head_rows[-1]})",
                   bold=True, top=True)
    row += 1
    if net:
        put("Net profit", f"=B{gp}-B{indirect}", bold=True, top=True)
    else:
        put("Total costs (direct + indirect)", f"=B{direct}+B{indirect}", bold=True, top=True)
        put("Sales left after all costs", f"=B{sales_row}-B{direct}-B{indirect}", bold=True)


AUTO_NOTE = ("Breakage / returns / transport and Compliance GST are calculated automatically "
             "as a percentage of COGS = product cost + labour (total direct costs). The "
             "percentages are set on Monthly inputs.")


def cost_split_sheet(ws: Worksheet, d: MonthData, largest_first: bool = False) -> None:
    row = title(ws, "Indirect vs direct cost %", d.label,
                ["Direct costs: product cost and labour on the month's invoices. "
                 "Indirect costs: entered on Monthly inputs.",
                 AUTO_NOTE])
    if not d.has_inputs:
        row = banner(ws, row, "No indirect costs were entered for this month.", 3)
    _statement(ws, row, d, net=False, largest_first=largest_first)
    _finish(ws)


def pnl_sheet(ws: Worksheet, d: MonthData, largest_first: bool = False) -> None:
    row = title(ws, "Profit & loss", d.label,
                ["Sales exclude GST (GST collected is not income). Invoices not included "
                 "(open issues) are not in these figures.",
                 AUTO_NOTE])
    if not d.has_inputs:
        row = banner(ws, row, "No indirect costs were entered for this month.", 3)
    _statement(ws, row, d, net=True, largest_first=largest_first)
    _finish(ws)


# ---------------------------------------------------------------------------
# 11 Payment modes
# ---------------------------------------------------------------------------
def payment_sheet(ws: Worksheet, d: MonthData, payments: list[Payment] | None) -> None:
    """
    Per invoice: amounts received by mode. Source per invoice: the payments
    export (every payment applied to that invoice, whatever its date);
    otherwise the payment mode printed on the invoice with 'Payment Made'.
    The rest of the invoice total is 'Not received'.
    """
    by_invoice: dict[str, list[Payment]] = {}
    for p in payments or []:
        by_invoice.setdefault(p.invoice_no, []).append(p)
    rows, modes, accounts = [], [], {}
    for i in sorted(d.invoices, key=lambda x: (x.invoice_date, x.invoice_no)):
        got: dict[str, float] = {}
        for p in by_invoice.get(i.invoice_no, []):
            got[p.mode] = got.get(p.mode, 0.0) + p.amount
            if p.deposit_to:
                accounts[p.deposit_to] = accounts.get(p.deposit_to, 0.0) + p.amount
        if not got and i.payment_mode and i.payment_made:
            got[i.payment_mode] = i.payment_made
        for m in got:
            if m not in modes:
                modes.append(m)
        rows.append(dict(invoice_no=i.invoice_no, date=i.invoice_date,
                         customer=i.customer, total=i.total,
                         **{f"mode:{m}": round(a, 2) for m, a in got.items()}))
    modes.sort()
    source = ("Zoho payments export" if payments is not None else
              "payment mode printed on the invoices (no payments export chosen)")
    row = title(ws, "Payment mode analysis", d.label,
                [f"Source: {source}. All payments applied to the month's invoices are "
                 "counted, whatever the payment date.",
                 "'Not received' = invoice total less the payments found for it."])

    mode_cols = [Col(m, f"mode:{m}", "money", 14, total="sum") for m in modes]
    detail_cols = ([Col("Invoice no", "invoice_no", width=17), Col("Date", "date", "date", 11),
                    Col("Customer", "customer", width=26),
                    Col("Invoice total", "total", "money", 14, total="sum")]
                   + mode_cols
                   + [Col("Not received", "pending", "money", 14,
                          formula="=MAX(0,{total}{r}-SUM(%s))" % (
                              "{first_mode}{r}:{last_mode}{r}" if modes else "0"),
                          total="sum")])
    # Work out exactly where each block goes, because the summary's formulas
    # must point at the detail table written after it:
    #   section title, heading, one row per mode + "Not received", total,
    #   2 blank rows; then the same for accounts; then the detail table.
    sum_head = row + 1
    after = sum_head + (len(modes) + 1) + 1 + 3
    acc_section = after
    if accounts:
        after = (after + 1) + len(accounts) + 1 + 3
    detail_head = after + 1
    first, last = detail_head + 1, detail_head + max(1, len(rows))
    letters = {c.key: get_column_letter(k) for k, c in enumerate(detail_cols, start=1)}
    if modes:
        for c in detail_cols:
            if c.formula:
                c.formula = c.formula.replace("{first_mode}", letters[f"mode:{modes[0]}"]) \
                                     .replace("{last_mode}", letters[f"mode:{modes[-1]}"])

    # Summary by mode (formulas over the detail table).
    row = section(ws, row, "By mode")
    total_col = letters["total"]
    summary = [dict(mode=m, col=letters[f"mode:{m}"]) for m in modes] + \
              [dict(mode="Not received", col=letters["pending"])]
    ws_rows = []
    for s in summary:
        ws_rows.append(dict(mode=s["mode"],
                            n=f'=COUNTIF({s["col"]}${first}:{s["col"]}${last},">0")',
                            amount=f'=SUM({s["col"]}${first}:{s["col"]}${last})'))
    t = table(ws, row, [
        Col("Mode", "mode", width=17),
        Col("Invoices", "n", "qty", 11),
        Col("Amount", "amount", "money", 14, total="sum"),
        Col("% of invoice total", "share", "pct", 12,
            formula=f'=IF(SUM({total_col}${first}:{total_col}${last})=0,"",'
                    f'{{amount}}{{r}}/SUM({total_col}${first}:{total_col}${last}))',
            total='=IF({amount}{r}=0,"",SUM({share}{first}:{share}{last}))'),
    ], ws_rows)

    if accounts:
        row = section(ws, acc_section, "By account deposited to (payments export)")
        t = table(ws, row, [Col("Deposited to", "acc", width=30),
                            Col("Amount", "amount", "money", 14, total="sum")],
                  [dict(acc=a, amount=round(v, 2)) for a, v in sorted(accounts.items())])
    row = section(ws, detail_head - 1, "Per invoice")
    table(ws, row, detail_cols, rows, filters=True)
    _finish(ws)


# ---------------------------------------------------------------------------
# Not included
# ---------------------------------------------------------------------------
def not_included_sheet(ws: Worksheet, d: MonthData) -> None:
    row = title(ws, "Not included", d.label, [
        "Invoices with open issues on Scan review are left out of every report. Fix them "
        "on Scan review and generate again to include them."])
    t = table(ws, row, [
        Col("Invoice no", "invoice_no", width=17),
        Col("Date", "date", "date", 11),
        Col("Customer", "customer", width=26),
        Col("Invoice total", "total", "money", 14, total="sum"),
        Col("Why", "why", width=90),
    ], [dict(invoice_no=x.invoice_no, date=x.invoice_date, customer=x.customer,
             total=x.total, why=" | ".join(x.reasons)) for x in d.left_out],
        empty_text="Every invoice read is included in the reports.")
    row = section(ws, t["next"], "Files skipped during the scan")
    table(ws, row, [Col("File", "file_name", width=34), Col("Reason", "reason", width=90)],
          [dict(file_name=f["file_name"], reason=f["reason"]) for f in d.skipped_files],
          empty_text="No files were skipped.")
    # Long reasons wrap inside their cell instead of running off the page.
    for r in ws.iter_rows(min_row=t["head"] + 1):
        for c in r:
            if c.column in (2, 5) and isinstance(c.value, str) and len(c.value) > 30:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    _finish(ws)