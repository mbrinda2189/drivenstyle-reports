"""
rto_list.py - Reading the dealership's monthly delivery list ("RTO list")
=========================================================================

WHAT THE FILE IS
----------------
Drive N Style works inside a Hyundai dealership. Every month the dealership
gives a list of the new cars delivered ("RTO List - SEP.xlsx"), one row per
car:

    S.NO | DELIVERY DATE | CUSTOMER NAME | MODEL | VIN NO | SALES CONSULTANT |
    Location | OE Accessories List | OE Total Value | DNS Accessories List |
    DNS Total Value | Executive Contribution | DNS Team Contribution |
    Executive Incentive | Remarks

    OE  = the car maker's own (original equipment) accessories
    DNS = Drive N Style accessories sold on that car

WHY THE TOOL READS IT (v0.9.0)
------------------------------
The invoices only show the cars that BOUGHT something. This list shows every
car delivered, so together they answer: of the cars delivered, how many took
DNS accessories (penetration), which did not and why (remarks), and does the
list agree with what was actually invoiced.

HOW A CAR IS LINKED TO ITS INVOICE
----------------------------------
Zoho's invoice carries the last six digits of the VIN - in the "VIN /
Registration Number" field and at the end of the customer name ("SOUDHA
BEGAM 271295"). The list has the full VIN (MALB281CYTM271295) or just the
last six digits. `vin6` is those six digits on both sides.

READING RULES
-------------
* The heading row is found by its headings (not by position), so extra
  title rows above it do not matter. Empty columns are ignored.
* A row without a customer name is skipped.
* Values are "-" or blank when there is nothing: read as 0 / empty.
* The delivery date is text like 01.09.26; a real Excel date also works.
* Model names are typed many ways ("I20", "I 20", "All New i20", "New
  Venue"): `model_key` reduces them to one name so they can be grouped.

This module has no database and no Qt code.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


class RtoFileError(Exception):
    """The file could not be read as a delivery list (plain-language text)."""


@dataclass
class RtoCar:
    row: int                    # row number in the sheet
    delivered: date | None
    customer: str
    model: str                  # as typed
    model_group: str            # standard name, see model_key
    vin: str
    vin6: str                   # last six digits, "" if not six digits
    consultant: str
    location: str               # upper case; "(not given)" if blank
    oe_list: str
    oe_value: float
    dns_list: str
    dns_value: float
    executive_contribution: str
    team_contribution: str
    incentive: float            # "Executive Incentive" column
    remarks: str

    @property
    def has_dns(self) -> bool:
        return self.dns_value > 0 or bool(self.dns_list)

    @property
    def has_oe(self) -> bool:
        return self.oe_value > 0 or bool(self.oe_list)


# heading (letters and digits only, lower case) -> field
HEADINGS = {
    "sno": "sno", "deliverydate": "delivered", "customername": "customer",
    "model": "model", "vinno": "vin", "vin": "vin",
    "salesconsultant": "consultant", "location": "location",
    "oeaccessorieslist": "oe_list", "oetotalvalue": "oe_value",
    "dnsaccessorieslist": "dns_list", "dnstotalvalue": "dns_value",
    "executivecontribution": "executive_contribution",
    "dnsteamcontribution": "team_contribution",
    "executiveincentive": "incentive", "remarks": "remarks",
}
REQUIRED = ("customer", "model", "vin", "dns_value")


def _key(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _text(value) -> str:
    text = " ".join(str(value if value is not None else "").split())
    return "" if text == "-" else text


def _amount(value) -> float:
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = re.sub(r"[^0-9.\-]", "", str(value or ""))
    try:
        return round(float(text), 2)
    except ValueError:
        return 0.0


def _day(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    m = re.match(r"\s*(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", str(value or ""))
    if not m:
        return None
    d, mo, y = (int(x) for x in m.groups())
    try:
        return date(y + 2000 if y < 100 else y, mo, d)
    except ValueError:
        return None


def vin6(text) -> str:
    """
    The six digits that identify the car: the end of a 17-character VIN, or
    a 6-digit number standing at the end of the text ("SOUDHA BEGAM 271295",
    "310221"). A registration number (TN45 BK5353) or a phone number gives
    "" - those are not delivered-car invoices.
    """
    text = str(text if text is not None else "").replace(" ", "").upper()
    if len(text) == 17 and text[-6:].isdigit():
        return text[-6:]
    m = re.search(r"(?<!\d)(\d{6})$", text)
    return m.group(1) if m else ""


def model_key(model: str) -> str:
    """
    One name per car model: "All New i20", "I 20", "I20" -> I20; "New Venue"
    -> VENUE; "Grand i10 Nios", "Nios" -> NIOS. Variants that are sold as
    different cars keep their own name (I20 N LINE, CRETA EV).
    """
    text = re.sub(r"\s+", " ", str(model or "").upper()).strip()
    text = re.sub(r"^(ALL NEW|NEW)\s+", "", text)
    text = re.sub(r"\bI\s+(\d+)", r"I\1", text)
    if "NIOS" in text:
        return "NIOS"
    return text or "(not given)"


def read_rto(path: str | Path) -> list[RtoCar]:
    """Read the delivery list; raises RtoFileError with a plain message."""
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        raise RtoFileError(f"“{Path(path).name}” could not be opened as an Excel "
                           f"file ({exc}).") from exc
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for h, heading_row in enumerate(rows[:15]):
            cols = {HEADINGS[_key(c)]: i for i, c in enumerate(heading_row)
                    if _key(c) in HEADINGS}
            if all(f in cols for f in REQUIRED):
                return _cars(rows[h + 1:], cols, h + 2)
    raise RtoFileError(
        f"“{Path(path).name}” does not look like the delivery (RTO) list: no row "
        "with the headings Customer Name, Model, VIN No and DNS Total Value was found.")


def _cars(rows: list[tuple], cols: dict[str, int], first_row: int) -> list[RtoCar]:
    def get(row, field):
        i = cols.get(field)
        return row[i] if i is not None and i < len(row) else None

    out = []
    for n, row in enumerate(rows, start=first_row):
        customer = _text(get(row, "customer"))
        if not customer:
            continue
        vin = _text(get(row, "vin"))
        model = _text(get(row, "model"))
        out.append(RtoCar(
            row=n, delivered=_day(get(row, "delivered")), customer=customer,
            model=model, model_group=model_key(model), vin=vin, vin6=vin6(vin),
            consultant=_text(get(row, "consultant")),
            location=_text(get(row, "location")).upper() or "(not given)",
            oe_list=_text(get(row, "oe_list")), oe_value=_amount(get(row, "oe_value")),
            dns_list=_text(get(row, "dns_list")), dns_value=_amount(get(row, "dns_value")),
            executive_contribution=_text(get(row, "executive_contribution")),
            team_contribution=_text(get(row, "team_contribution")),
            incentive=_amount(get(row, "incentive")), remarks=_text(get(row, "remarks"))))
    return out
