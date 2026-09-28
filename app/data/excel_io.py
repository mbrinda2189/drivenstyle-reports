"""
excel_io.py - Reading the client's master sheets and exporting masters
======================================================================

WHAT THIS MODULE DOES
---------------------
Import happens in three small steps, so the screen can show the user what
will happen before anything is saved:

    1. read_sheet(path, master)
           Opens an .xlsx or .csv file and finds the heading row. Client
           sheets often start with a title such as "Drive N Style - Product
           Master" or a blank line, so the first 15 rows are checked and the
           row whose cells best match the master's known headings is used.
           Returns the headings and the data rows below them.

    2. suggest_mapping(master, headings, remembered)
           Proposes which column holds which field. Column matches the user
           confirmed at the previous import come first; otherwise headings
           are compared with each field's name and synonyms ("Rate" ->
           Selling price, "Mobile" -> Phone). The user can change any match
           on screen.

    3. convert_rows(master, sheet, mapping)
           Turns each data row into clean values: amounts such as
           "₹ 2,400.00" into numbers, "Yes"/"Y"/"✓" into True, "Goods" into
           the category Product, HSN codes such as 8708.0 into "8708", and so
           on. Rows with a value that cannot be understood are listed as
           problems (with their sheet row number) and left out; blank rows
           are ignored.

export_rows(master, rows, path) writes a master to .xlsx with the same
headings the import recognises, so an exported file can be edited in Excel
and imported back.

export_audit(entries, path) writes audit-log entries (as returned by
MastersRepo.audit_entries) to .xlsx, for filing or review outside the tool.

A "Serial no" column (S.No) in the client's sheets is simply not matched to
any field, so it is ignored.

Old-style .xls files are not read: open them in Excel and "Save As" .xlsx.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.data.master_defs import FieldDef, MasterDef, header_key
from app.utils import parse_inr

SUPPORTED = (".xlsx", ".xlsm", ".csv")
HEADER_SEARCH_ROWS = 15

# Words accepted for yes / no fields (compared in lower case).
_YES = {"yes", "y", "true", "1", "active", "✓", "✔", "x", "applicable", "a"}
_NO = {"no", "n", "false", "0", "inactive", "not applicable", "na", "n/a",
       "-", "nil"}

# Extra words accepted for the product Category field.
_CATEGORY_WORDS = {
    "product": "Product", "products": "Product", "goods": "Product",
    "good": "Product", "item": "Product", "accessory": "Product",
    "accessories": "Product", "material": "Product",
    "service": "Service", "services": "Service", "labour": "Service",
    "labor": "Service", "job": "Service", "work": "Service",
}


class ImportFileError(Exception):
    """The file cannot be used (wrong type, empty, unreadable)."""


@dataclass
class SheetData:
    """Headings and data rows of one sheet."""
    headers: list[str]
    rows: list[tuple[int, list]]      # (row number in the sheet, cell values)
    sheet_names: list[str]            # all sheets in the file (for a chooser)
    sheet: str                        # the sheet that was read
    header_row: int                   # row number of the headings


# ---------------------------------------------------------------------------
# Step 1 - read the file
# ---------------------------------------------------------------------------
def read_sheet(path: str | Path, master: MasterDef,
               sheet: str | None = None) -> SheetData:
    """Read `sheet` (default: the first) of an .xlsx or .csv file."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xls":
        raise ImportFileError(
            "Old-style .xls files cannot be read. Open the file in Excel and "
            "use Save As → Excel Workbook (.xlsx), then import that file.")
    if ext not in SUPPORTED:
        raise ImportFileError("Choose an Excel (.xlsx) or CSV file.")

    try:
        if ext == ".csv":
            sheet_names, sheet = ["CSV"], "CSV"
            raw = _read_csv(path)
        else:
            wb = load_workbook(path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            sheet = sheet if sheet in sheet_names else sheet_names[0]
            raw = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
            wb.close()
    except ImportFileError:
        raise
    except Exception as exc:                   # damaged / locked file
        raise ImportFileError(f"The file could not be opened: {exc}") from exc

    header_index = _find_header_row(raw, master)
    if header_index is None:
        raise ImportFileError(
            f"No heading row was found in sheet “{sheet}”. The sheet should "
            "have one row of column headings above the data.")

    headers = _clean_headers(raw[header_index])
    width = len(headers)
    rows = []
    for i in range(header_index + 1, len(raw)):
        cells = (list(raw[i]) + [None] * width)[:width]
        if any(not _blank(c) for c in cells):
            rows.append((i + 1, cells))           # +1: sheet rows start at 1
    return SheetData(headers, rows, sheet_names, sheet, header_index + 1)


def _read_csv(path: Path) -> list[list]:
    """Read a CSV, trying UTF-8 first (Excel's 'CSV UTF-8') then Windows-1252."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with open(path, newline="", encoding=encoding) as fh:
                return [row for row in csv.reader(fh)]
        except UnicodeDecodeError:
            continue
    raise ImportFileError("The CSV file's text encoding was not recognised.")


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _find_header_row(raw: list[list], master: MasterDef) -> int | None:
    """
    Index of the heading row: among the first rows, the one with the most
    cells matching a known field heading. If none match, the first row with
    at least two filled cells is used (the user then matches columns by hand).
    """
    best, best_score = None, 0
    first_filled = None
    for i, row in enumerate(raw[:HEADER_SEARCH_ROWS]):
        texts = [c for c in row if isinstance(c, str) and c.strip()]
        if first_filled is None and sum(not _blank(c) for c in row) >= 2:
            first_filled = i
        score = sum(any(f.matches_header(t) for f in master.importable_fields)
                    for t in texts)
        if score > best_score:
            best, best_score = i, score
    return best if best is not None else first_filled


def _clean_headers(row: list) -> list[str]:
    """
    Turn the heading cells into unique, non-empty names. Blank headings
    become "Column C" (by Excel letter); repeated headings get " (2)".
    Trailing blank columns are dropped.
    """
    cells = list(row)
    while cells and _blank(cells[-1]):
        cells.pop()
    out, seen = [], {}
    for i, cell in enumerate(cells):
        name = str(cell).strip() if not _blank(cell) else f"Column {_letter(i)}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _letter(index: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA (Excel column letters)."""
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ---------------------------------------------------------------------------
# Step 2 - suggest which column is which field
# ---------------------------------------------------------------------------
def suggest_mapping(master: MasterDef, headers: list[str],
                    remembered: dict[str, str] | None = None
                    ) -> dict[str, str | None]:
    """
    field key -> column heading (or None if no column fits).
    Each column is used for at most one field.
    """
    remembered = remembered or {}
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    fields = master.importable_fields

    # 1. Matches the user confirmed last time, if that heading still exists.
    for f in fields:
        h = remembered.get(f.key)
        if h in headers and h not in used:
            mapping[f.key] = h
            used.add(h)

    # 2. Exact label match, then synonym match - in two passes so that a
    #    heading such as "Labour charge" goes to Labour charge, not to the
    #    looser "Labour" synonym of Labour involved.
    for strict in (True, False):
        for f in fields:
            if mapping.get(f.key):
                continue
            for h in headers:
                if h in used:
                    continue
                exact = header_key(h) in (header_key(f.label), header_key(f.key))
                if exact or (not strict and f.matches_header(h)):
                    mapping[f.key] = h
                    used.add(h)
                    break
    return {f.key: mapping.get(f.key) for f in fields}


# ---------------------------------------------------------------------------
# Step 3 - convert rows to clean values
# ---------------------------------------------------------------------------
def convert_rows(master: MasterDef, sheet: SheetData,
                 mapping: dict[str, str | None]
                 ) -> tuple[list[dict], list[str]]:
    """
    Returns (records, problems).

    records   one dict per usable row, holding ONLY the matched fields plus
              "_row" (the row number in the sheet, for messages)
    problems  plain-language messages for rows that were left out
    """
    index = {h: i for i, h in enumerate(sheet.headers)}
    matched = [(master.get_field(k), index[h])
               for k, h in mapping.items() if h and h in index]
    records, problems = [], []
    for row_no, cells in sheet.rows:
        if all(_blank(cells[i]) for _, i in matched):
            continue                               # nothing we use: ignore
        record, row_problems = {"_row": row_no}, []
        for f, i in matched:
            try:
                record[f.key] = convert_value(f, cells[i])
            except ValueError as exc:
                row_problems.append(f"{f.label} {exc}")
        if row_problems:
            problems.append(f"Row {row_no}: " + "; ".join(row_problems) + ".")
        else:
            records.append(record)
    return records, problems


def convert_value(f: FieldDef, value):
    """
    Convert one cell to the field's type. Raises ValueError with a short
    reason (e.g. '“abc” is not an amount') if it cannot.
    """
    if f.kind == "money":
        if _blank(value):
            return 0.0
        if isinstance(value, (int, float)):
            number = float(value)
        else:
            number = parse_inr(str(value))
            if number is None:
                raise ValueError(f"“{value}” is not an amount")
        if number < 0:
            raise ValueError("cannot be negative")
        return round(number, 2)

    if f.kind == "bool":
        if _blank(value):
            return bool(f.default)
        if isinstance(value, bool):
            return value
        word = _as_text(value).lower()
        if word in _YES:
            return True
        if word in _NO:
            return False
        raise ValueError(f"“{value}” should be Yes or No")

    if f.kind == "choice":
        if _blank(value):
            return None                 # decided later (default / inferred)
        word = _as_text(value)
        for choice in f.choices:
            if word.lower() == choice.lower():
                return choice
        if f.key == "category" and word.lower() in _CATEGORY_WORDS:
            return _CATEGORY_WORDS[word.lower()]
        if f.open_choice:
            return word
        raise ValueError(
            f"“{value}” should be one of {', '.join(f.choices)}")

    # "text" and "lookup" (e.g. Incentive group) are kept as clean text; a
    # lookup name is checked against the other master when it is saved.
    return _as_text(value)


def _as_text(value) -> str:
    """
    Cell value as clean text. Whole numbers lose the '.0' Excel adds, so
    HSN 8708.0 -> "8708" and phone 9876543210.0 -> "9876543210".
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%d-%m-%Y")
    return " ".join(str(value).split())


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="0B2545")   # theme navy


def export_rows(master: MasterDef, rows: list[dict], path: str | Path) -> None:
    """
    Write `rows` (as returned by MastersRepo.list_rows) to an .xlsx file:
    one heading row in the theme's navy, amounts as numbers with two
    decimals, yes/no as "Yes"/"No", dates as dd-mm-yyyy.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = master.title[:31]            # Excel limit on sheet names
    fields = list(master.fields)
    ws.append([f.label for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        out = []
        for f in fields:
            v = row.get(f.key)
            if f.kind == "bool":
                v = "Yes" if v else "No"
            elif f.kind == "date":
                v = v.strftime("%d-%m-%Y") if v else ""
            out.append(v)
        ws.append(out)

    for col, f in enumerate(fields, start=1):
        letter = _letter(col - 1)
        if f.kind == "money":
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "#,##0.00"
        ws.column_dimensions[letter].width = max(
            12, min(45, max((len(str(c.value or "")) for c in ws[letter]),
                            default=12) + 2))
    ws.freeze_panes = "A2"
    wb.save(path)


AUDIT_COLUMNS = (("at", "Date & time", 20), ("user", "User", 14),
                 ("master", "Master", 16), ("record", "Record", 32),
                 ("action", "Action", 12), ("field", "Field", 20),
                 ("old_value", "Old value", 40), ("new_value", "New value", 40),
                 ("source", "Source", 30))


def export_audit(entries: list[dict], path: str | Path,
                 master_titles: dict[str, str] | None = None) -> None:
    """
    Write audit-log entries to an .xlsx file: one row per entry, newest
    first, date & time as dd-mm-yyyy hh:mm:ss, master shown by its title
    (e.g. "Sales executives").
    """
    titles = master_titles or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit log"
    ws.append([label for _, label, _ in AUDIT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for e in entries:
        row = []
        for key, _, _ in AUDIT_COLUMNS:
            value = e.get(key, "")
            if key == "at" and value:
                value = datetime.fromisoformat(value).strftime("%d-%m-%Y %H:%M:%S")
            elif key == "master":
                value = titles.get(value, value)
            row.append(value)
        ws.append(row)
    for col, (_, _, width) in enumerate(AUDIT_COLUMNS):
        ws.column_dimensions[_letter(col)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)