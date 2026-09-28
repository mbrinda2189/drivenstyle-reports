"""
payments_io.py - Reading Zoho's "Payments Received" export
==========================================================

WHAT THIS MODULE DOES
---------------------
`read_payments(path)` reads the optional export chosen in step 3 of the
Generate page (CSV or Excel) and returns one `Payment` per row:

    invoice_no   "Invoice Number"             e.g. DNS26-GST-0749
    mode         "Mode"                       UPI / Cash / Bank Transfer ...
    amount       "Amount Applied to Invoice"  (falls back to "Amount")
    date         "Date"                       the payment date
    deposit_to   "Deposit To"                 e.g. Drive N Style - HDFC

WHY "AMOUNT APPLIED TO INVOICE"
-------------------------------
One invoice is often paid in parts by different modes (Rs. 25,000 UPI +
Rs. 50,000 bank transfer), and one payment row belongs to one invoice.
Report 11 adds up every payment applied to each of the month's invoices,
whatever the payment date - a September payment can settle an August
invoice, and it then counts for August.

Only the columns above are needed; others are ignored. Headings are
matched ignoring capitals and spaces. A file without "Invoice Number" and
"Mode" columns raises PaymentsFileError.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.data.master_defs import header_key
from app.utils import parse_inr


class PaymentsFileError(Exception):
    """The file is not a usable payments export."""


@dataclass
class Payment:
    invoice_no: str
    mode: str
    amount: float
    date: date | None = None
    deposit_to: str = ""


COLUMNS = {                      # our name -> accepted headings (header_key form)
    "invoice_no": ("invoicenumber", "invoiceno", "invoice"),
    "mode": ("mode", "paymentmode", "modeofpayment"),
    "applied": ("amountappliedtoinvoice",),
    "amount": ("amount", "amountreceived"),
    "date": ("date", "paymentdate"),
    "deposit_to": ("depositto",),
}


def read_payments(path: str | Path) -> list[Payment]:
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".csv":
            with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
                rows = [r for r in csv.reader(fh)]
        elif ext in (".xlsx", ".xlsm"):
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            wb.close()
        else:
            raise PaymentsFileError("Choose the Zoho payments export (.csv or .xlsx).")
    except PaymentsFileError:
        raise
    except Exception as exc:
        raise PaymentsFileError(f"The payments file could not be opened: {exc}") from exc
    if not rows:
        raise PaymentsFileError("The payments file is empty.")

    heads = [header_key(h or "") for h in rows[0]]
    col = {}
    for name, accepted in COLUMNS.items():
        for i, h in enumerate(heads):
            if h in accepted and name not in col:
                col[name] = i
    if "invoice_no" not in col or "mode" not in col:
        raise PaymentsFileError("This does not look like Zoho's Payments Received "
                                "export: the Invoice Number and Mode columns are missing.")

    out = []
    for r in rows[1:]:
        def cell(name):
            i = col.get(name)
            return r[i] if i is not None and i < len(r) else None
        inv = str(cell("invoice_no") or "").strip()
        if not inv:
            continue
        raw = cell("applied") if cell("applied") not in (None, "") else cell("amount")
        amount = raw if isinstance(raw, (int, float)) else parse_inr(str(raw or "")) or 0.0
        when = cell("date")
        if isinstance(when, datetime):
            when = when.date()
        elif isinstance(when, str) and when.strip():
            try:
                when = date.fromisoformat(when.strip()[:10])
            except ValueError:
                when = None
        out.append(Payment(inv, str(cell("mode") or "").strip() or "Not stated",
                           round(float(amount), 2), when if isinstance(when, date) else None,
                           str(cell("deposit_to") or "").strip()))
    return out