"""
test_web_monthly.py - The monthly reports tool through the web server (v0.26.0)
===============================================================================

The reading, matching, inputs and workbook are tested in test_invoice_export,
test_invoices, test_reports ... Here a small September (the two real
invoices of those tests plus the special cases) goes through the WEB
addresses from upload to download, to prove that:

    * staff cannot use any monthly address (cost and profit)
    * a file is checked when uploaded, and a wrong one does not replace
      the month's file
    * "Read invoices" gives the same counts and log as the desktop screen
    * every kind of Scan review fix works, is saved for the right scope
      and is in the audit log with the admin's e-mail
    * saved matches can be changed and removed
    * Monthly inputs are checked and saved
    * the workbook's figures are exactly those the desktop calculation
      gives for the same database (snapshot)
    * History: download, regenerate, remove (the files go too), remove month
"""

import io

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("itsdangerous")

from openpyxl import load_workbook                          # noqa: E402

from app.data.database import connect                       # noqa: E402
from app.data.inputs_repo import InputsRepo                 # noqa: E402
from app.data.invoices_repo import InvoicesRepo             # noqa: E402
from app.data.masters_repo import MastersRepo               # noqa: E402
from app.reports import snapshot as snap                    # noqa: E402
from app.reports import pdf_export                          # noqa: E402
from tests.test_invoice_export import extra_rows, rows_0753, rows_226, write_csv   # noqa: E402
from tests.test_web_auth import ADMIN, H, make, sign_in     # noqa: E402

STAFF = "asha@gmail.com"
M = "/api/monthly/2026-09"
ITEMS = [("I20 - PVC Full Floor Mat + Labour Extra", "Product", 1800, 150),
         ("Underbody Coating - 5 Seater", "Service", 900, 0),
         ("Silencer Coating - All Cars", "Service", 300, 0),
         ("Sunfilm - Nano Ceramic - Front (SK)", "Service", 2400, 300),
         ("Sunfilm - Nano Ceramic - Side and Rear (SK)", "Service", 3100, 500),
         ("Horn", "Product", 1100, 0), ("Punch.EV Mudflap - Techno", "Product", 300, 0),
         ("PUNCH - PVC Full Floor Mat + Labour Extra", "Product", 1900, 150),
         ("Punch EV Number Plate Frame", "Product", 200, 0)]


def put_file(client, kind, data: bytes, name: str, month=M):
    return client.post(f"{month}/files/{kind}", params={"filename": name}, content=data, headers=H)


@pytest.fixture
def world(tmp_path):
    app, settings = make(tmp_path)
    admin = sign_in(app, ADMIN)
    admin.post("/api/users", json={"email": STAFF}, headers=H)
    for name, category, cost, labour in ITEMS:
        made = admin.post("/api/masters/products/rows", headers=H, json={
            "effective_from": "2026-04-01",
            "values": {"name": name, "category": category, "cost_price": cost,
                       "has_labour": bool(labour), "labour_charge": labour}})
        assert made.status_code == 201, made.text
    admin.post("/api/masters/executives/rows", headers=H, json={
        "values": {"name": "Nandha Kumar", "phone": "9000000003", "branch": "HO"}})
    admin.post("/api/masters/cars/rows", headers=H, json={
        "values": {"make": "Hyundai", "model": "i20", "segment": "Hatchback"}})
    export = write_csv(tmp_path / "Invoice.csv", rows_0753() + rows_226() + extra_rows())
    return settings, admin, sign_in(app, STAFF), export.read_bytes()


def test_monthly_tool_is_for_admins_only(world):
    _, _, staff, export = world
    assert staff.get(f"{M}/overview").status_code == 403
    assert put_file(staff, "invoices", export, "Invoice.csv").status_code == 403
    for address in (f"{M}/read", f"{M}/generate", f"{M}/fix", "/api/monthly/runs/1/regenerate"):
        assert staff.post(address, json={}, headers=H).status_code == 403
    for address in (f"{M}/issues", f"{M}/invoices", f"{M}/matches", f"{M}/inputs",
                    "/api/monthly/history", "/api/monthly/choices",
                    "/api/monthly/runs/1/download", f"{M}/issues/export"):
        assert staff.get(address).status_code == 403
    assert staff.delete(f"{M}/invoices", headers=H).status_code == 403


def test_files_are_checked_when_uploaded(world):
    settings, admin, _, export = world
    start = admin.get(f"{M}/overview").json()
    assert start["label"] == "September 2026" and start["scan"] is None
    assert start["files"] == {"invoices": None, "payments": None, "rto": None}
    assert len(start["reports"]) == 12 and start["masters"]["products"] == 9
    assert admin.post(f"{M}/read", headers=H).status_code == 400          # nothing uploaded
    files = put_file(admin, "invoices", export, "C:\\Zoho\\Invoice.csv").json()
    assert files["invoices"]["name"] == "Invoice.csv" and files["invoices"]["by"] == ADMIN
    # The invoice export chosen as the payments export: refused, nothing kept.
    wrong = put_file(admin, "payments", export, "Invoice.csv")
    assert wrong.status_code == 400 and admin.get(f"{M}/overview").json()["files"]["payments"] is None
    assert put_file(admin, "invoices", b"not,an,export\n1,2,3\n", "x.csv").status_code == 400
    assert admin.get(f"{M}/overview").json()["files"]["invoices"]["name"] == "Invoice.csv"
    assert put_file(admin, "rto", b"x", "list.csv").status_code == 400     # .xlsx only
    assert put_file(admin, "salaries", b"x", "a.csv").status_code == 404
    assert put_file(admin, "invoices", export, "a.csv", "/api/monthly/2026-13").status_code == 404
    kept = sorted(p.name for p in (settings.data_dir / "months" / "2026-09").iterdir())
    assert kept == ["files.json", "invoices.csv"]
    assert admin.delete(f"{M}/files/invoices", headers=H).json()["invoices"] is None


def read_month(admin, export):
    put_file(admin, "invoices", export, "Invoice.csv")
    reply = admin.post(f"{M}/read", headers=H)
    assert reply.status_code == 200, reply.text
    return reply.json()


def test_read_invoices_like_the_desktop(world):
    _, admin, _, export = world
    done = read_month(admin, export)
    texts = [e["text"] for e in done["log"]]
    assert "3 invoices (17 lines in the file) read for September 2026." in texts
    assert "1 invoice dated in other months ignored." in texts
    assert any(t.startswith("DNS-300-2627:") for t in texts)               # the void invoice
    scan = done["scan"]
    assert (scan["read"], scan["skipped"]) == (3, 1) and scan["source"] == "Invoice.csv"
    assert scan["open_issues"] > 0 and texts[-1] == f"{scan['open_issues']} issues to review."
    # An export without the month: nothing is saved over the month already read.
    put_file(admin, "invoices", write_august(export), "August.csv")
    again = admin.post(f"{M}/read", headers=H)
    assert again.status_code == 400 and "Nothing was saved" in again.json()["detail"]
    assert admin.get(f"{M}/overview").json()["scan"]["read"] == 3


def write_august(export: bytes) -> bytes:
    """The same export with every September date moved to August."""
    return export.replace(b"2026-09-", b"2026-08-")


def test_scan_review_fixes_matches_and_the_audit_log(world):
    _, admin, _, export = world
    read_month(admin, export)
    found = admin.get(f"{M}/issues").json()["issues"]
    kinds = {(i["kind"], i["printed"]): i for i in found if i["status"] == "open"}
    edhayan = kinds[("salesperson", "Edhayan - Cuddalore")]
    seat = kinds[("product", "Seat Cover Lavish - Creta")]
    punch = kinds[("car", "PUNCH.EV")]
    totals = next(i for i in found if i["kind"] == "totals")
    choices = admin.get("/api/monthly/choices").json()
    assert choices["salesperson"][-1] == {"id": 0, "text": "Others (not in master)"}
    horn = next(c["id"] for c in choices["product"] if c["text"] == "Horn")

    # Nothing chosen, an id that does not exist, an issue that is not open: refused.
    assert admin.post(f"{M}/fix", json={"kind": "product", "key": seat["key"]},
                      headers=H).status_code == 400
    assert admin.post(f"{M}/fix", json={"kind": "car", "key": punch["key"], "target_id": 999},
                      headers=H).status_code == 400
    assert admin.post(f"{M}/fix", json={"kind": "product", "key": "no such item", "target_id": horn},
                      headers=H).status_code == 404

    done = admin.post(f"{M}/fix", json={"kind": "product", "key": seat["key"], "target_id": horn},
                      headers=H).json()
    assert done["done"] == "“Seat Cover Lavish - Creta” will be read as Horn."
    done = admin.post(f"{M}/fix", headers=H, json={
        "kind": "salesperson", "key": edhayan["key"], "target_id": 0}).json()
    assert done["done"] == "Others (not in master) set for all invoices showing “Edhayan - Cuddalore”."
    admin.post(f"{M}/fix", json={"kind": "car", "key": punch["key"], "target_id": 0}, headers=H)
    done = admin.post(f"{M}/fix", json={"kind": "totals", "key": totals["key"]}, headers=H).json()
    assert done["done"] == f"Invoice {totals['key']} accepted."
    left = {(i["kind"], i["printed"]) for i in done["issues"] if i["status"] == "open"}
    assert ("salesperson", "Edhayan - Cuddalore") not in left and ("car", "PUNCH.EV") not in left
    assert done["scan"]["open_issues"] == len(left) < len(kinds)

    log = admin.get("/api/audit", params={"master": "scan"}).json()["entries"]
    assert {e["user"] for e in log} == {ADMIN} and {e["source"] for e in log} == {"Scan review"}
    assert len(log) == 4

    # Saved matches: the four choices; change one, remove one.
    saved = admin.get(f"{M}/matches").json()
    assert sorted((m["store"], m["kind"]) for m in saved) == [
        ("ack", "totals"), ("alias", "car"), ("alias", "executive"), ("alias", "product")]
    item = next(m for m in saved if m["kind"] == "product")
    assert item["target"] == "Horn" and item["type_label"] == "Item" and item["invoices"] == 1
    pick = admin.get("/api/monthly/match-choices/product").json()
    mud = next(c["id"] for c in pick if c["text"].startswith("Punch.EV Mudflap"))
    body = {"store": "alias", "kind": "product", "key": item["key"]}
    after = admin.post(f"{M}/matches/change", json={**body, "target_id": mud}, headers=H).json()
    assert next(m for m in after if m["kind"] == "product")["target"] == "Punch.EV Mudflap - Techno"
    ack = next(m for m in saved if m["store"] == "ack")
    refuse = admin.post(f"{M}/matches/change", headers=H, json={
        "store": "ack", "kind": ack["kind"], "key": ack["key"], "target_id": 1})
    assert refuse.status_code == 400 and "can only be removed" in refuse.json()["detail"]
    after = admin.post(f"{M}/matches/remove", json=body, headers=H).json()
    assert not any(m["kind"] == "product" for m in after)
    assert admin.post(f"{M}/matches/remove", json=body, headers=H).status_code == 404
    back = admin.get(f"{M}/issues").json()["issues"]
    assert any(i["printed"] == "Seat Cover Lavish - Creta" and i["status"] == "open" for i in back)

    rows = admin.get(f"{M}/invoices").json()
    one = next(r for r in rows if r["invoice_no"] == "DNS-226-2627")
    assert one["executive"].startswith("Nandha Kumar") and one["gst"] == "No GST"
    assert one["items"] == 6 and len(one["lines"]) == 9 and one["total"] == 20000
    for address, name in ((f"{M}/issues/export", "Issues"), (f"{M}/matches/export", "Saved_matches")):
        sheet = admin.get(address)
        assert sheet.status_code == 200 and name in sheet.headers["content-disposition"]
        assert load_workbook(io.BytesIO(sheet.content)).sheetnames


def test_monthly_inputs(world):
    _, admin, _, _ = world
    start = admin.get(f"{M}/inputs").json()
    assert start["costs"] == [] and not start["saved"] and start["threshold"] == 40
    assert "Rent" in start["default_heads"]
    assert [r["pct"] for r in start["auto_rates"]] == [4.0, 3.0]
    twice = admin.put(f"{M}/inputs", headers=H, json={
        "costs": [["Rent", 45000], ["rent", 1]], "threshold": 40})
    assert twice.status_code == 400 and "entered twice" in twice.json()["detail"]
    assert admin.put(f"{M}/inputs", json={"costs": [], "threshold": 140}, headers=H).status_code == 400
    saved = admin.put(f"{M}/inputs", headers=H, json={
        "costs": [["Rent", 45000], ["Salaries", 120000], ["", 0]], "threshold": 35}).json()
    assert saved["costs"] == [["Rent", 45000.0], ["Salaries", 120000.0]] and saved["saved"]
    assert saved["threshold"] == 35
    october = admin.get("/api/monthly/2026-10/inputs").json()
    assert october["previous"] == {"label": "September 2026",
                                   "costs": [["Rent", 45000.0], ["Salaries", 120000.0]]}
    rates = admin.put("/api/monthly/auto-rates", json={"rates": {"auto_breakage_pct": 5}}, headers=H)
    assert [r["pct"] for r in rates.json()] == [5.0, 3.0]
    assert admin.put("/api/monthly/auto-rates", json={"rates": {"auto_breakage_pct": 500}},
                     headers=H).status_code == 400
    log = admin.get("/api/audit", params={"master": "inputs"}).json()["entries"]
    assert len(log) == 4 and {e["user"] for e in log} == {ADMIN}


def test_workbook_has_the_desktop_figures_and_history_works(world, monkeypatch):
    settings, admin, _, export = world
    monkeypatch.setattr(pdf_export, "find_libreoffice", lambda: None)      # no PDF here
    assert admin.post(f"{M}/generate", json={"reports": ["Profit & loss"]},
                      headers=H).status_code == 400                        # month not read yet
    read_month(admin, export)
    assert admin.post(f"{M}/generate", json={"reports": []}, headers=H).status_code == 400
    assert admin.post(f"{M}/generate", json={"reports": ["Balance sheet"]},
                      headers=H).status_code == 400
    made = admin.post(f"{M}/generate", headers=H, json={
        "reports": ["Profit & loss", "Invoice-wise profitability"], "pdf": True})
    assert made.status_code == 200, made.text
    out = made.json()
    run = out["run"]
    assert run["available"] and not run["pdf_available"] and "LibreOffice" in out["pdf_error"]
    assert run["file_name"].startswith("DriveNStyle_Sep-2026_Reports_") and run["user"] == ADMIN
    assert not out["payments_used"] and not out["rto_used"]

    # The figures are the desktop calculation's, for the very same database.
    conn = connect(settings.db_path)
    masters = MastersRepo(conn, "check")
    figures = snap.snapshot(masters, InvoicesRepo(masters), InputsRepo(masters), 2026, 9)
    conn.close()
    assert run["sales"] == figures["totals"]["sales"]
    assert run["gross_profit"] == figures["totals"]["gross_profit"]
    assert (run["invoices"], run["left_out"]) == (figures["totals"]["invoices"],
                                                  figures["totals"]["left_out"])

    book = admin.get(f"/api/monthly/runs/{run['id']}/download")
    assert book.status_code == 200 and run["file_name"] in book.headers["content-disposition"]
    names = load_workbook(io.BytesIO(book.content)).sheetnames
    assert "Not included" in names and any("Profit" in n for n in names)
    assert admin.get(f"/api/monthly/runs/{run['id']}/download", params={"pdf": True}).status_code == 404
    assert admin.post(f"/api/monthly/runs/{run['id']}/pdf", headers=H).status_code == 400

    history = admin.get("/api/monthly/history").json()
    assert [r["id"] for r in history["runs"]] == [run["id"]] and not history["pdf_possible"]
    assert history["months"][0]["label"] == "September 2026" and history["months"][0]["invoices"] == 3
    assert admin.get(f"{M}/overview").json()["last_run"]["id"] == run["id"]

    again = admin.post(f"/api/monthly/runs/{run['id']}/regenerate", headers=H).json()["run"]
    assert again["id"] != run["id"] and again["reports"] == ["Invoice-wise profitability",
                                                             "Profit & loss"]
    assert again["sales"] == run["sales"]
    assert admin.get("/api/monthly/runs/999/download").status_code == 404

    output = settings.data_dir / "output"
    assert len(list(output.glob("*.xlsx"))) >= 1
    assert admin.delete(f"{M}/runs", headers=H).json()["removed"] == 2
    assert list(output.glob("*.xlsx")) == [] and admin.get("/api/monthly/history").json()["runs"] == []
    assert admin.delete(f"{M}/invoices", headers=H).json() == {"removed": 3}
    assert admin.get(f"{M}/overview").json()["scan"] is None
