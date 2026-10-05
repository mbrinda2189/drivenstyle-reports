"""
test_packages.py - Package sales recognised from the invoice (v0.8.0)
=====================================================================

An invoice is a package sale when every item of a package (Packages
master) is on it. The package incentive then replaces the items' own
incentives; an invoice missing exactly one item is "almost a package".
"""

from datetime import date

import openpyxl
import pytest

from app.data.database import connect
from app.data.inputs_repo import InputsRepo
from app.data.masters_repo import MastersRepo, RowChange
from app.reports.data import build_month
from app.reports.generate import generate
from tests.test_counter_sales import month
from tests.test_invoice_export import line

ITEMS = {"Graphene Coating": ("Graphene Coating - Turtle Wax", 27000, "Graphene Coating"),
         "Underbody Coating": ("Underbody Coating - 5 Seater", 3500, "Underbody"),
         "Silencer Coating": ("Silencer Coating - All Cars", 1500, ""),
         "Full Exterior Glass Coating": ("Glass Coating - All Cars", 1000, "")}


@pytest.fixture
def masters():
    m = MastersRepo(connect(":memory:"))
    m.save("incentives", [
        RowChange(None, dict(name="Graphene Protection Package", incentive_amount=1850,
                             bill_value=25000, active=True), date(2026, 4, 1)),
        RowChange(None, dict(name="Graphene Coating", incentive_amount=1500,
                             bill_value=25000, active=True), date(2026, 4, 1)),
        RowChange(None, dict(name="Underbody", incentive_amount=200,
                             bill_value=3500, active=True), date(2026, 4, 1))])
    m.save("executives", [RowChange(None, dict(
        name="Mano Vikram", phone="9876543210", branch="HO", active=True))])
    m.save("cars", [RowChange(None, dict(make="Hyundai", model="Creta",
                                         segment="SUV", active=True))])
    for name, price, group in list(ITEMS.values()) + [("Horn", 2100, "")]:
        m.save("products", [RowChange(None, dict(
            sku="", name=name, hsn_sac="", category="Product", incentive_group=group,
            selling_price=price, cost_price=price / 2, has_labour=False,
            labour_charge=0, active=True, vehicle_needed=True), date(2026, 4, 1))])
    r = m.import_records("package_items", [
        dict(package="Graphene Protection Package", item=item, product=zoho, _row=n)
        for n, (item, (zoho, _, _)) in enumerate(ITEMS.items(), start=2)])
    assert r.added == 4 and not r.skipped
    return m


def rows(no, names, total):
    return [line(no, "2026-09-21", "Mano Vikram", "Creta", n, "", "", "goods", 1, p, p,
                 inv_total=total) for n, p in names]


def test_package_needs_every_item(tmp_path, masters):
    full = [(z, p) for z, p, _ in ITEMS.values()] + [("Horn", 2100)]
    irepo = month(tmp_path, masters, rows("DNS-1-2627", full, 35100)
                  + rows("DNS-2-2627", full[:3], 32000)            # no glass coating
                  + rows("DNS-3-2627", full[1:3], 5000))           # underbody + silencer
    assert irepo.issues(2026, 9) == []
    d = build_month(masters, irepo, InputsRepo(masters), 2026, 9)
    one, two, three = sorted(d.invoices, key=lambda i: i.invoice_no)

    assert one.package.package == "Graphene Protection Package"
    assert len(one.package.lines) == 4                 # the horn is a normal sale
    assert one.package.list_value == 33000 and one.package.coupon_value == 25000
    assert one.package.incentive == 1850
    # the package incentive replaces the items' own incentives
    assert all(l.is_package and not l.incentive_group for l in one.package.lines)

    assert two.package is None
    assert (two.near_package, two.near_missing) == (
        "Graphene Protection Package", "Full Exterior Glass Coating")
    assert [l.incentive_group for l in two.lines][:2] == ["Graphene Coating", "Underbody"]
    assert three.package is None and not three.near_package


def test_package_item_must_be_a_product(masters):
    r = masters.import_records("package_items", [dict(
        package="Graphene Protection Package", item="Teflon", product="No such item",
        _row=2)])
    assert r.added == 0 and "not in the Product master" in r.skipped[0]


def test_workbook_package_sheet_and_incentive(tmp_path, masters):
    full = [(z, p) for z, p, _ in ITEMS.values()]
    irepo = month(tmp_path, masters, rows("DNS-1-2627", full, 33000))
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, InputsRepo(masters), 2026, 9, out,
                 ["Basic package analysis", "Spot incentive calculation"])
    wb = openpyxl.load_workbook(r.path)
    cells = [c for row in wb["5 Packages"].iter_rows(values_only=True) for c in row]
    assert "Graphene Protection Package" in cells and "DNS-1-2627" in cells
    inc = [row for row in wb["7 Spot incentive"].iter_rows(values_only=True)
           if row[1] == "DNS-1-2627"]
    assert len(inc) == 1                               # one row: the package
    assert inc[0][3] == "Package (4 items)" and inc[0][6] == 1850 and inc[0][7] == 25000


def test_automatic_indirect_costs(tmp_path, masters):
    """v0.8.1: 4% and 3% of COGS (product cost + labour) every month; a typed
    head with the same meaning is not counted twice."""
    full = [(z, p) for z, p, _ in ITEMS.values()]
    irepo = month(tmp_path, masters, rows("DNS-1-2627", full, 33000))
    inputs = InputsRepo(masters)
    inputs.save_costs(2026, 9, [("Rent", 1000.0), ("Breakage and returns", 500.0)])
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert d.cogs == 16500                                  # half of 33,000, no labour
    assert d.auto_indirect == [("Breakage / returns / transport", 0.04, 660.0),
                               ("Compliance GST", 0.03, 495.0)]
    assert d.entered_indirect == [("Rent", 1000.0)]
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, inputs, 2026, 9, out, ["Profit & loss"])
    col = {row[0]: row[1] for row in
           openpyxl.load_workbook(r.path)["12 Profit & loss"].iter_rows(values_only=True)}
    assert col["Rent"] == 1000 and "Breakage and returns" not in col
    assert col["Breakage / returns / transport"].endswith("*0.04,2)")
    assert col["Compliance GST"].endswith("*0.03,2)")


def test_automatic_indirect_percentages_can_be_changed(tmp_path, masters):
    """v0.8.2: one setting for all months, logged in the audit log."""
    full = [(z, p) for z, p, _ in ITEMS.values()]
    irepo = month(tmp_path, masters, rows("DNS-1-2627", full, 33000))
    inputs = InputsRepo(masters)
    assert inputs.auto_rates() == {"auto_breakage_pct": 4.0, "auto_compliance_pct": 3.0}
    inputs.save_auto_rates({"auto_breakage_pct": 5.0, "auto_compliance_pct": 3.0})
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert d.auto_indirect == [("Breakage / returns / transport", 0.05, 825.0),
                               ("Compliance GST", 0.03, 495.0)]
    log = masters.conn.execute(
        "SELECT field, old_value, new_value FROM audit_log WHERE master = 'inputs'"
    ).fetchall()
    assert [tuple(r) for r in log] == [
        ("Breakage / returns / transport (% of product cost)", "4%", "5%")]
    with pytest.raises(ValueError):
        inputs.save_auto_rates({"auto_breakage_pct": 150})


def test_executive_marked_no_incentive(tmp_path, masters):
    """v0.11.0: "Gets incentive = No" - the sale counts, the incentive does not."""
    full = [(z, p) for z, p, _ in ITEMS.values()]
    irepo = month(tmp_path, masters, rows("DNS-1-2627", full, 33000)        # a package
                  + rows("DNS-2-2627", full[:2], 30500))                    # two items
    inputs = InputsRepo(masters)
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert masters.list_rows("executives")[0]["gets_incentive"] is True     # default Yes
    assert sum(i.incentive_payable for i in d.invoices) > 0 and d.no_incentive == []

    ex = masters.list_rows("executives")[0]
    masters.save("executives", [RowChange(ex["id"], {**ex, "gets_incentive": False})])
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert sum(i.incentive_payable for i in d.invoices) == 0
    assert d.no_incentive == ["Mano Vikram – HO"]
    assert d.sales == 63500 and all(i.package or len(i.lines) == 2 for i in d.invoices)
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, inputs, 2026, 9, out, ["Spot incentive calculation"])
    ws = openpyxl.load_workbook(r.path)["7 Spot incentive"]
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert any("Gets incentive = No" in c and "Mano Vikram" in c for c in cells)
    assert "DNS-1-2627" not in cells and "DNS-2-2627" not in cells          # no rows at all
    log = masters.conn.execute("SELECT field, old_value, new_value FROM audit_log "
                               "WHERE field = 'Gets incentive'").fetchone()
    assert tuple(log) == ("Gets incentive", "Yes", "No")


def test_internal_team_incentive(tmp_path, masters):
    """v0.12.0: Rs. 3,000 per car PPF to the internal team, on top of the
    salesperson's incentive; nothing at all for a two-wheeler PPF."""
    from app.data.database import apply_internal_ppf_rule
    masters.save("incentives", [RowChange(None, dict(
        name="PPF", incentive_amount=3000, bill_value=70000, active=True), date(2026, 4, 1))])
    for name, price in (("Paint Protection Film - 8 Year - Blaupunkt", 85000),
                        ("Paint Protection Film-5years-Turtle wax-twowheeler", 2000),
                        ("Paint Protection Film -Labour Charges", 1)):
        masters.save("products", [RowChange(None, dict(
            sku="", name=name, hsn_sac="", category="Product", incentive_group="PPF",
            selling_price=price, cost_price=price / 2, has_labour=False, labour_charge=0,
            active=True, vehicle_needed=True), date(2026, 4, 1))])
    apply_internal_ppf_rule(masters.conn)
    masters.conn.commit()
    got = {p["name"]: (p["internal_incentive"], p["incentive_group"])
           for p in masters.list_rows("products") if "Paint" in p["name"]}
    assert got["Paint Protection Film - 8 Year - Blaupunkt"] == (3000, "PPF")
    assert got["Paint Protection Film-5years-Turtle wax-twowheeler"] == (0, "")   # nothing
    assert got["Paint Protection Film -Labour Charges"][0] == 0                   # Rs. 1 item

    irepo = month(tmp_path, masters,
                  rows("DNS-1-2627", [("Paint Protection Film - 8 Year - Blaupunkt", 85000)],
                       85000)
                  + rows("DNS-2-2627", [("Paint Protection Film-5years-Turtle wax-twowheeler",
                                         2000)], 2000))
    inputs = InputsRepo(masters)
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert d.internal_incentive == 3000
    car = next(i for i in d.invoices if i.invoice_no == "DNS-1-2627")
    assert car.internal_incentive == 3000 and car.incentive_payable == 3000   # both are paid
    bike = next(i for i in d.invoices if i.invoice_no == "DNS-2-2627")
    assert bike.internal_incentive == 0 and bike.incentive_payable == 0

    # the internal incentive does not depend on the salesperson getting incentive
    ex = masters.list_rows("executives")[0]
    masters.save("executives", [RowChange(ex["id"], {**ex, "gets_incentive": False})])
    d = build_month(masters, irepo, inputs, 2026, 9)
    assert d.internal_incentive == 3000 and sum(i.incentive_payable for i in d.invoices) == 0

    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, inputs, 2026, 9, out, ["Spot incentive calculation"])
    sheet = list(openpyxl.load_workbook(r.path)["7 Spot incentive"].iter_rows(values_only=True))
    team = [row for row in sheet if row[0] == "Internal team"]
    assert len(team) == 2                              # the summary line + one detail row
    detail = next(row for row in team if row[1] == "DNS-1-2627")
    assert detail[6] == 3000 and detail[7] == detail[8]            # paid in full
    from app.reports.pdf_book import incentive_by_executive
    last = incentive_by_executive(d)[-1]
    assert (last["executive"], last["payable"]) == ("Internal team", 3000)


def test_incentive_is_rounded_up_to_ten(tmp_path, masters):
    """v0.12.1: an executive's incentive for the month is rounded UP to the
    next Rs. 10 (1,492 -> 1,500; 2,677.80 -> 2,680), once, on the total."""
    from app.reports.data import round_up_10
    from app.reports.pdf_book import incentive_by_executive
    got = [round_up_10(x) for x in (1492, 2677.80, 1500, 0, 50.01, 1499.999)]
    assert got == [1500, 2680, 1500, 0, 60, 1500]
    # graphene at a discount: 1,500 x 21,000 / 25,000 = 1,260; underbody in full = 200;
    # a second invoice: underbody 200 x 3,333 / 3,500 = 190.46  ->  1,650.46 -> 1,660
    graphene, underbody = "Graphene Coating - Turtle Wax", "Underbody Coating - 5 Seater"
    data = [line("DNS-1-2627", "2026-09-21", "Mano Vikram", "Creta", graphene, "", "", "goods",
                 1, 27000, 21000, inv_total=24500),
            line("DNS-1-2627", "2026-09-21", "Mano Vikram", "Creta", underbody, "", "", "goods",
                 1, 3500, 3500, inv_total=24500),
            line("DNS-2-2627", "2026-09-22", "Mano Vikram", "Creta", underbody, "", "", "goods",
                 1, 3500, 3333, inv_total=3333)]
    irepo = month(tmp_path, masters, data)
    inputs = InputsRepo(masters)
    d = build_month(masters, irepo, inputs, 2026, 9)
    exact = sum(i.incentive_payable for i in d.invoices)
    assert exact == pytest.approx(1650.46, abs=0.01)
    assert d.incentive_by_executive == {"Mano Vikram – HO": 1660.0}
    assert d.incentive_payable == 1660.0
    assert incentive_by_executive(d)[0]["payable"] == 1660.0
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, inputs, 2026, 9, out,
                 ["Spot incentive calculation", "Executive-wise sales"])
    wb = openpyxl.load_workbook(r.path)
    summary = next(row for row in wb["7 Spot incentive"].iter_rows(values_only=True)
                   if row[0] == "Mano Vikram – HO" and str(row[3]).startswith("=CEILING"))
    assert summary[3].endswith(",2),10)")
    sales = next(row for row in wb["8 Executive-wise sales"].iter_rows(values_only=True)
                 if row[0] == "Mano Vikram – HO")
    assert str(sales[-1]).startswith("=CEILING(ROUND(SUMIF(")
