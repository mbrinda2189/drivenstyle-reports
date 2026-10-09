"""
test_web_imports.py - Masters from and to Excel through the web server (v0.25.0)
================================================================================

Reading a sheet, matching columns and saving rows are tested in
test_excel_io.py and test_masters_repo.py. Here it is proved that the web
addresses run that same code in the same order as the desktop dialog, and
the points that exist only on the web:

    * nothing is saved until Import; the preview shows what will be saved
    * a required column that is not matched blocks the import
    * the column matches are remembered for the next file
    * the uploaded file belongs to its uploader, and is gone after the
      import or Cancel
    * Zoho's item list: Rs. 1 labour items left out, products not in the
      list are reported, and ONLY AN ADMIN can remove them - exactly those
    * export gives the rows asked for, in the desktop layout
"""

import io

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("itsdangerous")

from openpyxl import Workbook, load_workbook           # noqa: E402

from tests.test_web_auth import ADMIN, H, make, sign_in    # noqa: E402

STAFF = "asha@gmail.com"


def xlsx(rows, title="Sheet1", more_sheets=()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    for name, sheet_rows in more_sheets:
        extra = wb.create_sheet(name)
        for row in sheet_rows:
            extra.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def send(client, master, data: bytes, name="items.xlsx"):
    return client.post(f"/api/masters/{master}/import/upload", params={"filename": name},
                       content=data, headers=H)


PRODUCTS = xlsx([
    ["Drive N Style - price list"],                       # a title above the headings
    ["S.No", "Item Name", "Rate", "Cost", "Labour Charges"],
    [1, "I20 - PVC Full Floor Mat", 3500, 1800, 150],
    [2, "Horn", "1,200", 700, "-"],
    [3, "Seat Cover", "mrp less 10%", 2000, 0],          # a price in words: left out
])

ZOHO = ("Item ID,Item Name,SKU,HSN/SAC,Rate,Purchase Rate,Product Type\n"
        "1,I20 - PVC Full Floor Mat,EA-01,39181090,3600,1850,goods\n"
        "2,Teflon Coating,TC-01,998729,2500,400,service\n"
        "3,Labour Charges for Sunfilm - Front,,998729,1,0,service\n").encode()


@pytest.fixture
def world(tmp_path):
    app, settings = make(tmp_path)
    admin = sign_in(app, ADMIN)
    admin.post("/api/users", json={"email": STAFF}, headers=H)
    return settings, admin, sign_in(app, STAFF)


def test_preview_shows_what_will_be_saved_and_saves_nothing(world):
    settings, _, staff = world
    reply = send(staff, "products", PRODUCTS, "C:\\Users\\me\\Price list.xlsx")
    assert reply.status_code == 200, reply.text
    s = reply.json()
    assert s["file_name"] == "Price list.xlsx" and s["header_row"] == 2 and s["data_rows"] == 3
    assert s["mapping"]["name"] == "Item Name" and s["mapping"]["selling_price"] == "Rate"
    assert s["mapping"]["cost_price"] == "Cost" and s["mapping"]["labour_charge"] == "Labour Charges"
    assert s["mapping"]["sku"] is None and s["missing"] == []
    assert s["count"] == 2 and len(s["problems"]) == 1 and "Row 5" in s["problems"][0]
    assert s["preview"][1] == {"name": "Horn", "selling_price": 1200.0, "cost_price": 700.0,
                               "labour_charge": 0.0}
    assert s["has_rates"] and s["default_date"].endswith("-04-01")      # first load: 1 April
    assert s["can_choose_add_new"] and not s["is_zoho"]
    assert staff.get("/api/masters").json()[0]["rows"] == 0             # nothing saved yet
    assert len(list((settings.data_dir / "uploads").iterdir())) == 2    # the file + its note


def test_unmatched_required_column_blocks_the_import(world):
    _, _, staff = world
    s = send(staff, "products", PRODUCTS).json()
    mapping = {**s["mapping"], "name": None}
    again = staff.post(f"/api/masters/products/import/{s['token']}/preview",
                       json={"mapping": mapping}, headers=H).json()
    assert again["missing"] == ["Product name"]
    run = staff.post(f"/api/masters/products/import/{s['token']}/run",
                     json={"mapping": mapping}, headers=H)
    assert run.status_code == 400 and "Choose the column for: Product name" in run.json()["detail"]
    assert staff.get("/api/masters").json()[0]["rows"] == 0


def test_import_saves_logs_remembers_and_forgets_the_file(world):
    settings, admin, staff = world
    s = send(staff, "products", PRODUCTS).json()
    # The person corrects one match: "S.No" is offered as the SKU.
    mapping = {**s["mapping"], "sku": "S.No"}
    done = staff.post(f"/api/masters/products/import/{s['token']}/run",
                      json={"mapping": mapping, "effective_from": "2026-04-01"}, headers=H)
    assert done.status_code == 200, done.text
    r = done.json()
    assert (r["added"], r["updated"], r["unchanged"]) == (2, 0, 0)
    assert len(r["skipped"]) == 1 and r["not_in_zoho"] == []
    rows = staff.get("/api/masters/products/rows").json()["rows"]
    mat = next(x for x in rows if x["name"].startswith("I20"))
    assert mat["labour_charge"] == 150 and mat["has_labour"] is True      # labour above zero
    assert mat["effective_from"] == "2026-04-01" and mat["sku"] == "1"
    log = admin.get("/api/audit", params={"master": "products"}).json()["entries"]
    assert {(e["user"], e["source"]) for e in log} == {(STAFF, "Import: items.xlsx")}
    assert list((settings.data_dir / "uploads").iterdir()) == []          # file not kept
    assert staff.post(f"/api/masters/products/import/{s['token']}/preview", json={},
                      headers=H).status_code == 404
    # Next file: last time's match is suggested, and amounts now default to next month.
    s2 = send(staff, "products", PRODUCTS).json()
    assert s2["mapping"]["sku"] == "S.No" and not s2["default_date"].endswith("-04-01")
    # The same sheet again changes nothing.
    r2 = staff.post(f"/api/masters/products/import/{s2['token']}/run",
                    json={"mapping": s2["mapping"], "effective_from": "2026-04-01"},
                    headers=H).json()
    assert (r2["added"], r2["updated"], r2["unchanged"]) == (0, 0, 2)


def test_add_new_unticked_only_updates(world):
    _, _, staff = world
    staff.post("/api/masters/products/rows", json={"values": {"name": "Horn", "category": "Product"},
                                                   "effective_from": "2026-04-01"}, headers=H)
    s = send(staff, "products", PRODUCTS).json()
    r = staff.post(f"/api/masters/products/import/{s['token']}/run",
                   json={"mapping": s["mapping"], "effective_from": "2026-04-01",
                         "add_new": False}, headers=H).json()
    assert (r["added"], r["updated"]) == (0, 1)
    assert any("“Add items…” is unticked" in w for w in r["warnings"])
    assert staff.get("/api/masters").json()[0]["rows"] == 1


def test_an_upload_belongs_to_its_uploader_and_cancel_removes_it(world):
    settings, admin, staff = world
    s = send(staff, "products", PRODUCTS).json()
    assert admin.post(f"/api/masters/products/import/{s['token']}/preview", json={},
                      headers=H).status_code == 404
    assert staff.post(f"/api/masters/cars/import/{s['token']}/preview", json={},
                      headers=H).status_code == 404               # another master's address
    assert staff.post("/api/masters/products/import/..%2F..%2Fx/preview", json={},
                      headers=H).status_code == 404
    assert staff.delete(f"/api/masters/products/import/{s['token']}", headers=H).json() == {"ok": True}
    assert list((settings.data_dir / "uploads").iterdir()) == []


def test_files_that_cannot_be_used_are_refused_and_not_kept(world):
    settings, _, staff = world
    assert "Save As" in send(staff, "cars", b"old", "cars.xls").json()["detail"]
    assert send(staff, "cars", b"x", "cars.pdf").status_code == 400
    assert send(staff, "cars", b"", "cars.xlsx").status_code == 400
    assert send(staff, "cars", b"not really excel", "cars.xlsx").status_code == 400
    assert send(staff, "salaries", PRODUCTS).status_code == 404
    assert not any((settings.data_dir / "uploads").glob("*"))


def test_sheet_chooser(world):
    _, _, staff = world
    book = xlsx([["Notes"], ["nothing useful here"]], title="Cover", more_sheets=[
        ("Executive ph.no", [["S.NO", "NAME", "CONTACT NO", "BRANCH"],
                             [1, "Edhayan", "9000000004", "Ooty"]])])
    first = send(staff, "executives", book)
    assert first.status_code == 400 and "No heading row" in first.json()["detail"]
    book = xlsx([["NAME", "CONTACT NO"], ["Kumaran", "9000000001"]], title="HO", more_sheets=[
        ("Ooty", [["NAME", "CONTACT NO", "BRANCH"], ["Edhayan", "9000000004", "Ooty"]])])
    s = send(staff, "executives", book).json()
    assert s["sheet_names"] == ["HO", "Ooty"] and s["preview"][0]["name"] == "Kumaran"
    assert s["default_date"] is None and not s["can_choose_add_new"]
    ooty = staff.post(f"/api/masters/executives/import/{s['token']}/preview",
                      json={"sheet": "Ooty"}, headers=H).json()
    assert ooty["sheet"] == "Ooty" and ooty["preview"][0]["branch"] == "Ooty"
    r = staff.post(f"/api/masters/executives/import/{s['token']}/run",
                   json={"sheet": "Ooty", "mapping": ooty["mapping"]}, headers=H).json()
    assert r["added"] == 1 and r["effective_from"] is None


def test_zoho_item_list_and_only_an_admin_removes_the_rest(world):
    _, admin, staff = world
    for name in ("I20 - PVC Full Floor Mat", "Old Item Not In Zoho"):
        staff.post("/api/masters/products/rows",
                   json={"values": {"name": name, "category": "Product", "selling_price": 100},
                         "effective_from": "2026-04-01"}, headers=H)
    # Staff may run the import, and are told what is not in Zoho's list ...
    s = send(staff, "products", ZOHO, "Item.csv").json()
    assert s["is_zoho"] and s["count"] == 2 and not s["can_choose_add_new"]
    assert s["mapping"]["category"] == "Product Type" and s["default_date"].endswith("-04-01")
    r = staff.post(f"/api/masters/products/import/{s['token']}/run",
                   json={"mapping": s["mapping"], "effective_from": "2026-04-01"},
                   headers=H).json()
    assert (r["added"], r["updated"]) == (1, 1) and r["may_remove"] is False
    assert [p["name"] for p in r["not_in_zoho"]] == ["Old Item Not In Zoho"]
    assert any("₹1 labour item" in w for w in r["warnings"])
    # ... but cannot remove them, and an admin cannot use the staff member's upload.
    assert staff.post(f"/api/masters/products/import/{s['token']}/remove-missing",
                      headers=H).status_code == 403
    assert admin.post(f"/api/masters/products/import/{s['token']}/remove-missing",
                      headers=H).status_code == 404
    rows = {x["name"]: x for x in staff.get("/api/masters/products/rows").json()["rows"]}
    assert len(rows) == 3 and rows["I20 - PVC Full Floor Mat"]["selling_price"] == 3600
    assert rows["Teflon Coating"]["category"] == "Service"
    # The admin imports the list and removes exactly the products it found.
    s = send(admin, "products", ZOHO, "Item.csv").json()
    r = admin.post(f"/api/masters/products/import/{s['token']}/run",
                   json={"mapping": s["mapping"], "effective_from": "2026-04-01"},
                   headers=H).json()
    assert r["may_remove"] is True and (r["added"], r["updated"], r["unchanged"]) == (0, 0, 2)
    assert admin.post(f"/api/masters/products/import/{s['token']}/run",
                      json={"mapping": s["mapping"]}, headers=H).status_code == 400   # once only
    gone = admin.post(f"/api/masters/products/import/{s['token']}/remove-missing", headers=H).json()
    assert gone == {"removed": 1, "names": ["Old Item Not In Zoho"]}
    assert admin.get("/api/masters").json()[0]["rows"] == 2
    last = admin.get("/api/audit", params={"action": "Deleted"}).json()["entries"][0]
    assert last["source"] == "Import: Item.csv (not in Zoho's item list)" and last["user"] == ADMIN
    assert admin.post(f"/api/masters/products/import/{s['token']}/remove-missing",
                      headers=H).status_code == 404                # the token is used up


def test_export_gives_the_rows_asked_for(world):
    _, _, staff = world
    ids = [staff.post("/api/masters/cars/rows", json={"values": {"make": "Hyundai", "model": m}},
                      headers=H).json()["id"] for m in ("i20", "Creta", "Verna")]
    some = staff.post("/api/masters/cars/export", json={"ids": ids[:2]}, headers=H)
    assert some.status_code == 200 and some.headers["x-file-name"].startswith("Cars_")
    sheet = load_workbook(io.BytesIO(some.content)).active
    assert [c.value for c in sheet[1]][:3] == ["Make", "Model", "Segment"]
    assert sorted(r[1] for r in sheet.iter_rows(min_row=2, values_only=True)) == ["Creta", "i20"]
    every = staff.post("/api/masters/cars/export", json={}, headers=H)
    assert load_workbook(io.BytesIO(every.content)).active.max_row == 4
