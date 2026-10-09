"""
test_pdf.py - The workbook as one PDF (v0.9.1)
==============================================

Excel makes the PDF (through PowerShell), so the real export can only run on
a Windows PC with Excel. Here: the page set-up written into the workbook,
what happens when the PDF cannot be made, and the command handed to Windows.
"""

import subprocess
import sys
from types import SimpleNamespace

import openpyxl
from pathlib import Path

import pytest

from app.data.inputs_repo import InputsRepo
from app.reports import pdf_export
from app.reports.generate import generate
from app.reports.pdf_export import PdfError, export_pdf
from tests.test_rto import world  # noqa: F401  (fixture)


def test_workbook_is_set_up_for_printing(world):  # noqa: F811
    masters, irepo, tmp_path = world
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, InputsRepo(masters), 2026, 9, out,
                 ["Invoice-wise profitability", "Profit & loss"])
    wb = openpyxl.load_workbook(r.path)
    for ws in wb:
        assert ws.page_setup.orientation == "landscape"
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert "&P" in ws.oddFooter.center.text
    # the long sheet repeats its heading row on every page
    assert wb["1 Invoice profitability"].print_title_rows


@pytest.mark.skipif(sys.platform == "win32", reason="needs a PC without Excel automation")
def test_pdf_problem_does_not_lose_the_workbook(world, monkeypatch):  # noqa: F811
    masters, irepo, tmp_path = world
    out = tmp_path / "out"
    out.mkdir()
    # v0.26.0: without Excel the PDF is made by LibreOffice - here neither is there.
    monkeypatch.setattr(pdf_export, "find_libreoffice", lambda: None)
    r = generate(masters, irepo, InputsRepo(masters), 2026, 9, out, ["Profit & loss"],
                 pdf=True)
    assert r.path.exists() and r.pdf_path is None
    assert "Windows PC with Microsoft Excel" in r.pdf_error


@pytest.mark.skipif(sys.platform == "win32", reason="LibreOffice is used where Excel is not")
def test_libreoffice_makes_the_pdf_where_there_is_no_excel(tmp_path, monkeypatch):
    """v0.26.0 (web server): LibreOffice converts in a folder of its own and
    the PDF is moved to where it was asked for."""
    book = tmp_path / "Reports Sep.xlsx"
    openpyxl.Workbook().save(book)
    target = tmp_path / "saved" / "Final name.pdf"
    seen = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"], seen["env"] = cmd, kwargs["env"]
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        (outdir / "Reports Sep.pdf").write_bytes(b"%PDF")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(pdf_export, "find_libreoffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(pdf_export.subprocess, "run", fake_run)
    assert export_pdf(book, target) == target.resolve() and target.read_bytes() == b"%PDF"
    assert seen["cmd"][0] == "/usr/bin/soffice" and "--headless" in seen["cmd"]
    assert seen["cmd"][-1] == str(book.resolve())
    assert any(a.startswith("-env:UserInstallation=file://") for a in seen["cmd"])
    assert seen["env"]["LANG"] == "en_IN.UTF-8"        # 13,80,298.18, not 1,380,298.18

    monkeypatch.setattr(pdf_export.subprocess, "run",
                        lambda cmd, **k: SimpleNamespace(returncode=1, stdout="", stderr="boom"))
    with pytest.raises(PdfError, match="LibreOffice could not make the PDF. \\(boom\\)"):
        export_pdf(book, target)


def test_excel_is_asked_for_one_pdf(tmp_path, monkeypatch):
    book = tmp_path / "Reports with space.xlsx"
    openpyxl.Workbook().save(book)
    seen = {}

    def fake_run(cmd, env, **kwargs):
        seen["cmd"], seen["env"] = cmd, env
        (tmp_path / "Reports with space.pdf").write_bytes(b"%PDF")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(subprocess, "run", fake_run)
    pdf = export_pdf(book)
    assert pdf.name == "Reports with space.pdf" and pdf.exists()
    assert seen["cmd"][0] == "powershell" and "ExportAsFixedFormat" in seen["cmd"][-1]
    assert seen["env"]["DNS_XLSX"].endswith("Reports with space.xlsx")

    # Excel missing -> a message an accountant can act on
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: SimpleNamespace(
        returncode=1, stdout="ERROR: Retrieving the COM class factory ... 80040154 "
                             "Class not registered", stderr=""))
    (tmp_path / "Reports with space.pdf").unlink()
    with pytest.raises(PdfError, match="Microsoft Excel was not found"):
        export_pdf(book)
    with pytest.raises(PdfError, match="no longer"):
        export_pdf(tmp_path / "missing.xlsx")
    assert pdf_export.pdf_path_for("a/b.xlsx").name == "b.pdf"


# --- v0.10.0: the PDF version (summary tables + graphs) and the trend --------------
def test_pdf_version_layout(world):  # noqa: F811
    """
    The PDF version: one flowing sheet, no graphs, page numbers (v0.10.3);
    from v0.13.0 the client's order - Profit & loss first - with no
    descriptions under the headings and only the sections on his note.
    """
    from app.data.rto_list import read_rto
    from app.reports.data import build_month
    from app.reports.pdf_book import OTHER_COL, incentive_by_executive, shift_formula, \
        write_pdf_workbook
    from tests.test_rto import write_list
    masters, irepo, tmp_path = world
    inputs = InputsRepo(masters)
    inputs.save_costs(2026, 9, [("Postage", 5.0), ("Rent", 5000.0), ("Salaries", 900.0)])
    d = build_month(masters, irepo, inputs, 2026, 9)
    rto = read_rto(write_list(tmp_path / "RTO.xlsx"))
    wb = openpyxl.load_workbook(write_pdf_workbook(d, [d, d], None, tmp_path / "pdf.xlsx",
                                                   "tester", rto=rto))
    assert wb.sheetnames == ["Report"]                 # one sheet = sections follow on
    ws = wb["Report"]
    assert not ws._charts                              # no graphs
    # v0.19.1: a page break is only ever placed BEFORE a heading, so a heading is
    # never left at the foot of a page with its table on the next one
    for brk in ws.row_breaks.brk:
        top = ws.cell(brk.id + 1, 1)
        assert top.value and top.font.b and top.font.sz in (11, 14), top.value
    text = [row[0] for row in ws.iter_rows(values_only=True) if isinstance(row[0], str)]
    wanted = ["Profit & loss", "New-car business", "Top 10 products by gross profit",
              "Top 10 services by gross profit", "Penetration by location",
              "By location - OE accessories", "By model - OE accessories",
              "Branches by gross profit", "Packages vs sales", "Labour calculation",
              "Vehicle-wise average per car", "Spot incentive calculation",
              "Payment mode analysis"]
    order = [text.index(t) for t in wanted]
    assert order == sorted(order) and text[0] == "Profit & loss"      # P&L on page 1
    for gone in ("Headline figures", "Drive N Style – Executive summary",
                 "Points needing attention", "Service vs product profitability",
                 "Indirect vs direct cost %", "New-car vs other business",
                 "Trend analysis (month on month)", "By model - DNS accessories"):
        assert gone not in text
    # no descriptions under the headings, no "Prepared on" line
    assert not any(c.value and c.font.sz == 9 for row in ws.iter_rows(max_col=1) for c in row)
    assert not any(t.startswith(("Prepared on", "Labour cost =", "Each invoice is one car"))
                   for t in text)
    # New-car business: exactly the five lines asked for
    at = text.index("New-car business")
    assert text[at + 2:at + 8] == [
        "Particulars", "Cars delivered", "Cars fitted with DNS accessories",
        "DNS penetration %", "DNS value as per list", "DNS value per car delivered"]
    # page set-up
    assert "&P" in ws.oddFooter.right.text and "Drive N Style" in ws.oddFooter.left.text
    # v0.19.2: a fixed print size (not "fit to width"), so page ends are exact
    assert not ws.sheet_properties.pageSetUpPr.fitToPage
    assert 80 <= int(ws.page_setup.scale) <= 90
    assert ws.max_column <= 8
    assert all((ws.column_dimensions[c].width or 0) <= OTHER_COL for c in "BCDEFGH")
    # formulas moved with their tables (also past the removed description rows):
    # "% of sales" divides by the Sales row of its own table
    for r in ws.iter_rows():
        if r[0].value == "Sales (excluding GST)" and isinstance(r[2].value, str):
            n = r[0].row
            assert r[2].value == f'=IF($B${n}=0,"",B{n}/$B${n})'
            assert ws.cell(n + 2, 3).value == f'=IF($B${n}=0,"",B{n + 2}/$B${n})'
    for row in ws.iter_rows():                          # no formula points at text / blanks
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and "SUM(" not in c.value:
                import re
                for col, r in re.findall(r"\$?([A-Z]{1,3})\$?(\d+)", c.value):
                    v = ws[f"{col}{r}"].value
                    assert v is not None and (not isinstance(v, str) or v.startswith("="))
    assert shift_formula('=IF($B$7=0,"",B9/$B$7)', 100) == '=IF($B$107=0,"",B109/$B$107)'
    assert shift_formula("=SUM(C8:C11)", 10) == "=SUM(C18:C21)"
    # indirect cost heads largest first (PDF only)
    heads = text[text.index("Indirect costs") + 1:text.index("Total indirect costs")]
    assert heads[:2] == ["Rent", "Salaries"]
    payable = [g["payable"] for g in incentive_by_executive(d)]
    assert payable == sorted(payable, reverse=True)

    # pieces that must stay together: heading + table; a sub-heading straight after
    # the section heading stays with it
    from app.reports.pdf_book import PAGE_POINTS, PAGE_USE, _blocks, _page_breaks, \
        print_scale
    oe = next(c.row for row in ws.iter_rows(max_col=1) for c in row
              if c.value == "OE accessories")
    end = next(c.row for row in ws.iter_rows(max_col=1) for c in row
               if c.value == "Branches by gross profit") - 3
    pieces = _blocks(ws, oe, end)
    assert len(pieces) == 2 and pieces[0][0] == oe
    assert ws.cell(pieces[1][0], 1).value == "By model - OE accessories"
    # a piece that does not fit the rest of the page starts the next page
    scale = print_scale(ws)
    assert scale == int(ws.page_setup.scale) and 84 <= scale <= 88
    # a page holds about 40 rows at that size (it was wrongly taken as 32)
    rows_per_page = int(PAGE_POINTS / (scale / 100) * PAGE_USE // 15)
    assert 38 <= rows_per_page <= 42
    sheet = openpyxl.Workbook().active
    # a piece that does not fit the rest of the page starts the next page ...
    _page_breaks(sheet, [(1, 20), (23, 23 + rows_per_page - 15), (60, 62)], scale)
    assert [b.id for b in sheet.row_breaks.brk][:1] == [22]
    # ... two short pieces after a full page share the next page (v0.19.2: "By
    # account deposited to" must not go to a page of its own) ...
    sheet = openpyxl.Workbook().active
    _page_breaks(sheet, [(1, rows_per_page), (rows_per_page + 3, rows_per_page + 17),
                         (rows_per_page + 20, rows_per_page + 27)], scale)
    assert [b.id for b in sheet.row_breaks.brk] == [rows_per_page + 2]
    # ... and a table longer than a page is cut where the page is full
    sheet = openpyxl.Workbook().active
    _page_breaks(sheet, [(1, 5), (8, 8 + rows_per_page + 9)], scale)
    assert [b.id for b in sheet.row_breaks.brk] == [7, 7 + rows_per_page]

    # without the delivery list the new-car sections are simply left out
    ws = openpyxl.load_workbook(write_pdf_workbook(d, [d], None, tmp_path / "p2.xlsx"))["Report"]
    text = [row[0] for row in ws.iter_rows(values_only=True) if isinstance(row[0], str)]
    assert "New-car business" not in text and "Penetration by location" not in text
    assert text[0] == "Profit & loss" and "Branches by gross profit" in text

    # the Excel workbook is unchanged: summary, notes and gross-profit ranking stay
    r = generate(masters, irepo, inputs, 2026, 9, tmp_path, ["Profit & loss"])
    excel = [c[0] for c in openpyxl.load_workbook(r.path)["Summary"].iter_rows(
        values_only=True) if c[0]]
    assert "Points needing attention" in excel and "Headline figures" in excel


def test_trend_sheet_and_saved_delivery_totals(world):  # noqa: F811
    from tests.test_rto import write_list
    masters, irepo, tmp_path = world
    inputs = InputsRepo(masters)
    out = tmp_path / "out"
    out.mkdir()
    r = generate(masters, irepo, inputs, 2026, 9, out, ["Trend analysis (month on month)"],
                 rto_path=str(write_list(tmp_path / "RTO.xlsx")))
    saved = inputs.rto_months()["2026-09"]
    assert (saved["cars"], saved["took_dns"]) == (4, 3)
    ws = openpyxl.load_workbook(r.path)["4 Trend"]
    col = {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row[0]}
    assert col["Sales (excluding GST)"] == 6300 and col["Invoices"] == 3
    assert col["Cars delivered (delivery list)"] == 4 and col["DNS penetration %"] == 0.75
    assert "Sales by branch" in col and "Top 10 products by sales" in col
    assert len(ws._charts) == 2                      # lines + penetration bars
