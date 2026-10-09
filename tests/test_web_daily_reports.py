"""
test_web_daily_reports.py - Payout slip, summary, month check, Excel (v0.28.0)
==============================================================================

The same sample September as test_web_daily.py, posted through the web
addresses. What is proved:

    * the slip groups as the desktop slip does and equals the register
    * the summary's totals are the register's (cancelled lines left out)
    * THE MONTH CHECK: for the same invoices the register and the monthly
      tool give the same labour, incentive and internal team amounts; the
      month-end rounding is shown apart and is not a difference; an invoice
      not scanned daily, or left out by the monthly tool, is named
    * the Excel export holds every line and invoice
    * all four are read-only and open to staff
"""

import io
from datetime import date

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("itsdangerous")

from openpyxl import load_workbook                              # noqa: E402

from tests.test_invoice_export import rows_0753, rows_226, write_csv    # noqa: E402
from tests.test_web_auth import H                               # noqa: E402
from tests.test_web_daily import TODAY, lines, scan, september, world  # noqa: E402,F401

M = "/api/monthly/2026-09"


def monthly_september(admin, tmp_path, rows):
    """Read a September export into the monthly tool (as Generate reports does)."""
    export = write_csv(tmp_path / "Invoice.csv", rows)
    put = admin.post(f"{M}/files/invoices", params={"filename": "Invoice.csv"},
                     content=export.read_bytes(), headers=H)
    assert put.status_code == 200, put.text
    read = admin.post(f"{M}/read", headers=H)
    assert read.status_code == 200, read.text
    return read.json()


def as_in_the_pdfs(rows, **changes):
    for row in rows:
        row.update(changes)
    return rows


def test_slip_and_summary_equal_the_register(world):                # noqa: F811
    _, _, staff, drawer = world
    september(drawer, staff)
    total = scan(staff)["posted_total"]
    to_pay = staff.get("/api/daily/slip")
    assert to_pay.status_code == 200 and "To pay - pending as on" in to_pay.text
    for block in ("Labour - Floor mat", "Labour - Sunfilm", "Edhayan", "Grand total"):
        assert block in to_pay.text
    assert to_pay.text.index("Labour - Floor mat") < to_pay.text.index("Edhayan")
    assert "window.print()" in to_pay.text

    mat, inc = "DNS26-GST-0753-LABM", "DNS26-GST-0753-INC"
    staff.post("/api/daily/pay", headers=H, json={"line_ids": [mat, inc], "paid_date": TODAY,
                                                  "mode": "Cash", "reference": "V-1"})
    paid = staff.get("/api/daily/slip", params={"paid_on": TODAY}).text
    assert f"Paid on {date.today():%d-%m-%Y}" in paid and "DNS26-GST-0753" in paid
    assert "DNS-226-2627" not in paid
    assert staff.get("/api/daily/slip", params={"paid_on": "today"}).status_code == 400

    staff.post("/api/daily/cancel", json={"invoice_no": "DNS-226-2627", "reason": "test"}, headers=H)
    now = lines(staff)
    live = [l for l in now.values() if l["status"] != "Cancelled"]
    s = staff.get("/api/daily/summary").json()
    assert round(sum(r["pending"] + r["paid"] for r in s["by_payee"]), 2) == round(
        sum(l["amount"] for l in live), 2) < total
    assert s["by_day"] == [{"date": TODAY, "lines": 2,
                            "amount": round(now[mat]["amount"] + now[inc]["amount"], 2)}]
    month = s["by_month"][0]
    assert month["label"] == "September 2026"
    assert month["due"] == round(month["paid"] + month["pending"] + month["hold"], 2)
    assert month["paid"] == s["by_day"][0]["amount"]


def test_month_check_agrees_with_the_monthly_tool(world, tmp_path):   # noqa: F811
    _, admin, staff, drawer = world
    september(drawer, staff)                       # 0753, 226 and 237 scanned daily
    scan(staff)
    # Before the monthly tool has read the month: the register's side only.
    alone = staff.get("/api/daily/month-check", params={"month_text": "2026-09"}).json()
    assert alone["monthly_read"] is False and alone["register_invoices"] == 3
    assert all(r["monthly"] is None for r in alone["rows"]) and alone["rounding"] is None

    # The monthly tool reads an export with 0753 and 226 only (237 is not in it).
    monthly_september(admin, tmp_path,
                      as_in_the_pdfs(rows_0753(), **{"Sales person": "Edhayan - Ooty"}) + rows_226())
    check = staff.get("/api/daily/month-check", params={"month_text": "2026-09"}).json()
    assert check["monthly_read"] and check["label"] == "September 2026"
    assert check["not_scanned"] == [] and check["left_out"] == []
    assert check["not_in_monthly"] == ["DNS-237-2627"]            # scanned daily, not in the export
    # For the invoices both tools hold, every figure is the same to the paisa ...
    by_label = {r["label"]: r for r in check["rows"]}
    posted = lines(staff)
    for line_id in ("DNS26-GST-0753-INC", "DNS-226-2627-INC"):
        payee = posted[line_id]["payee"]
        assert by_label[payee]["difference"] == 0 and by_label[payee]["monthly"] > 0
    sunfilm = by_label["Labour - Sunfilm"]
    assert sunfilm["register"] == sunfilm["monthly"] == 800 and sunfilm["difference"] == 0
    # ... and the only differences are the invoice the export does not have.
    extra = {l["type"] if l["type"].startswith("Labour") else l["payee"]: l["amount"]
             for l in posted.values() if l["invoice_no"] == "DNS-237-2627"}
    for r in check["rows"]:
        assert r["difference"] == round(extra.get(r["label"], 0.0), 2), r
    assert check["totals"]["difference"] == round(sum(extra.values()), 2)
    assert len(check["rows"]) >= 4
    # Month-end rounding is shown apart: each executive up to the next Rs. 10.
    rounding = check["rounding"]
    assert rounding["rounded"] >= rounding["exact"] and rounding["amount"] == round(
        rounding["rounded"] - rounding["exact"], 2)
    assert all(p["rounded"] % 10 == 0 and 0 <= p["rounded"] - p["exact"] < 10
               for p in rounding["by_payee"])

    # An invoice the monthly tool has but nobody scanned daily is named.
    admin.post("/api/daily/clear", json={"confirm": "CLEAR"}, headers=H)
    only = [inv for inv in drawer.invoices.values() if inv.invoice_no == "DNS-226-2627"][0]
    drawer.put(staff, only)
    scan(staff)
    check = staff.get("/api/daily/month-check", params={"month_text": "2026-09"}).json()
    assert [n["invoice_no"] for n in check["not_scanned"]] == ["DNS26-GST-0753"]
    assert check["not_scanned"][0]["why"] == "Not scanned" and check["totals"]["difference"] < 0
    assert staff.get("/api/daily/month-check", params={"month_text": "2026-13"}).status_code == 400


def test_excel_export_and_sign_in(world):                           # noqa: F811
    _, _, staff, drawer = world
    september(drawer, staff)
    scan(staff)
    book = staff.get("/api/daily/export")
    assert book.status_code == 200 and "Payout_register" in book.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(book.content))
    assert wb.sheetnames == ["Payouts", "Invoices"]
    assert wb["Payouts"].max_row == len(lines(staff)) + 1 and wb["Invoices"].max_row == 4
    first = [c.value for c in wb["Payouts"][2]]
    assert first[0].startswith("DNS") and first[2].year == 2026 and first[9] == "Pending"
    from fastapi.testclient import TestClient
    nobody = TestClient(staff.app)
    for address in ("/api/daily/slip", "/api/daily/summary", "/api/daily/export",
                    "/api/daily/month-check?month_text=2026-09"):
        assert nobody.get(address).status_code == 401
