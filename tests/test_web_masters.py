"""
test_web_masters.py - Masters and audit log through the web server (v0.24.0)
============================================================================

The rules themselves are tested in test_masters_repo.py. Here it is proved
that the web addresses reach those same rules, and who may use them:

    * staff can read and change the masters (Brinda, 09-10-2026); nobody
      can without signing in; the audit log is for admins only
    * a refusal comes back as the repository's own plain sentence
    * a changed amount is stored from the date sent, the earlier amount
      stays for earlier dates
    * every change is in the audit log with the e-mail of who made it
    * Delete all needs the word DELETE
    * "bring across" copies the desktop database, keeps the users, never
      changes the source and refuses to overwrite masters by mistake
"""

import io
from datetime import date

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("itsdangerous")

from fastapi.testclient import TestClient              # noqa: E402
from openpyxl import load_workbook                     # noqa: E402

from app.data.database import connect                  # noqa: E402
from app.data.masters_repo import MastersRepo, RowChange   # noqa: E402
from tests.test_web_auth import ADMIN, H, make, sign_in    # noqa: E402
from web.backend.bring_across import BringAcrossError, bring_across   # noqa: E402

STAFF = "asha@gmail.com"
MAT = {"name": "I20 - PVC Full Floor Mat", "category": "Product", "selling_price": 3500,
       "cost_price": 1800, "has_labour": True, "labour_charge": 150}


@pytest.fixture
def world(tmp_path):
    app, settings = make(tmp_path)
    admin = sign_in(app, ADMIN)
    admin.post("/api/users", json={"email": STAFF}, headers=H)
    return app, settings, admin, sign_in(app, STAFF)


def test_the_screens_get_the_masters_shape(world):
    _, _, _, staff = world
    masters = staff.get("/api/masters").json()
    assert [m["key"] for m in masters] == ["products", "executives", "cars",
                                           "incentives", "package_items"]
    products = masters[0]
    assert products["has_rates"] and products["filter_field"] == "category"
    cost = next(f for f in products["fields"] if f["key"] == "cost_price")
    assert cost["kind"] == "money" and cost["dated"] and cost["label"] == "Cost price (₹)"
    assert products["rows"] == 0


def test_needs_a_sign_in_and_a_known_master(world):
    app, _, _, staff = world
    nobody = TestClient(app)
    assert nobody.get("/api/masters").status_code == 401
    assert nobody.get("/api/masters/products/rows").status_code == 401
    assert nobody.post("/api/masters/products/rows", json={"values": MAT},
                       headers=H).status_code == 401
    assert staff.get("/api/masters/salaries/rows").status_code == 404


def test_staff_add_and_change_with_dated_amounts(world):
    _, _, _, staff = world
    made = staff.post("/api/masters/products/rows",
                      json={"values": MAT, "effective_from": "2026-04-01"}, headers=H)
    assert made.status_code == 201, made.text
    row = made.json()
    assert row["cost_price"] == 1800 and row["effective_from"] == "2026-04-01"
    assert row["vehicle_needed"] is True and row["active"] is True      # defaults
    # Cost goes up from 1 October; only that field is sent.
    changed = staff.put(f"/api/masters/products/rows/{row['id']}",
                        json={"values": {"cost_price": 2000}, "effective_from": "2026-10-01"},
                        headers=H).json()
    assert changed["cost_price"] == 2000 and changed["selling_price"] == 3500
    history = staff.get(f"/api/masters/products/rows/{row['id']}/rates").json()
    assert [(h["effective_from"], h["cost_price"]) for h in history] == [
        ("2026-10-01", 2000), ("2026-04-01", 1800)]
    # The wrong date is taken out again; the last one cannot be.
    left = staff.post(f"/api/masters/products/rows/{row['id']}/rates/delete",
                      json={"effective_from": "2026-10-01"}, headers=H).json()
    assert [h["cost_price"] for h in left] == [1800]
    last = staff.post(f"/api/masters/products/rows/{row['id']}/rates/delete",
                      json={"effective_from": "2026-04-01"}, headers=H)
    assert last.status_code == 400 and "only set of amounts" in last.json()["detail"]


def test_refusals_are_the_repositorys_own_sentences(world):
    _, _, _, staff = world
    staff.post("/api/masters/products/rows", json={"values": MAT}, headers=H)
    twice = staff.post("/api/masters/products/rows",
                       json={"values": {**MAT, "name": "i20 -  pvc full floor mat"}}, headers=H)
    assert twice.status_code == 400 and "Duplicates are not allowed" in twice.json()["detail"]
    phone = staff.post("/api/masters/executives/rows",
                       json={"values": {"name": "Edhayan", "phone": "98765"}}, headers=H)
    assert "must have 10 digits" in phone.json()["detail"]
    several = staff.post("/api/masters/products/rows",
                         json={"values": {"name": "", "category": "Goods", "cost_price": -5}},
                         headers=H).json()["detail"].split("\n")
    assert len(several) == 3
    group = staff.post("/api/masters/products/rows",
                       json={"values": {**MAT, "name": "PPF", "incentive_group": "PPF"}},
                       headers=H)
    assert "is not in the Incentives master" in group.json()["detail"]
    assert staff.put("/api/masters/products/rows/999", json={"values": {"sku": "X"}},
                     headers=H).status_code == 404
    assert staff.post("/api/masters/products/rows",
                      json={"values": {**MAT, "name": "Horn"}, "effective_from": "1 April"},
                      headers=H).status_code == 400


def test_lookup_names_inactive_and_delete(world):
    _, _, _, staff = world
    staff.post("/api/masters/incentives/rows",
               json={"values": {"name": "PPF", "incentive_amount": 1000, "bill_value": 60000}},
               headers=H)
    ppf = staff.post("/api/masters/products/rows",
                     json={"values": {**MAT, "name": "PPF - Sedan", "incentive_group": "PPF"}},
                     headers=H).json()
    assert staff.get("/api/masters/products/rows").json()["lookups"] == {"incentives": ["PPF"]}
    assert staff.post("/api/masters/products/active", json={"ids": [ppf["id"]], "active": False},
                      headers=H).json() == {"changed": 1}
    assert staff.get("/api/masters/products/rows").json()["rows"][0]["active"] is False
    assert staff.post("/api/masters/products/delete-all", json={"confirm": "delete"},
                      headers=H).status_code == 400
    assert staff.get("/api/masters").json()[0]["rows"] == 1
    assert staff.post("/api/masters/products/delete-all", json={"confirm": "DELETE"},
                      headers=H).json() == {"deleted": 1}
    assert staff.post("/api/masters/incentives/delete", json={"ids": [1]},
                      headers=H).json() == {"deleted": 1}


def test_audit_log_names_who_and_is_for_admins(world):
    _, _, admin, staff = world
    row = staff.post("/api/masters/cars/rows",
                     json={"values": {"make": "Hyundai", "model": "i20", "segment": "Hatchback"}},
                     headers=H).json()
    admin.put(f"/api/masters/cars/rows/{row['id']}", json={"values": {"segment": "Premium hatch"}},
              headers=H)
    assert staff.get("/api/audit").status_code == 403
    assert staff.get("/api/audit/export").status_code == 403
    log = admin.get("/api/audit", params={"master": "cars"}).json()
    assert [(e["user"], e["action"], e["field"], e["new_value"]) for e in log["entries"]] == [
        (ADMIN, "Edited", "Segment", "Premium hatch"),
        (STAFF, "Added", "", "Make: Hyundai; Model: i20; Segment: Hatchback; Active: Yes")]
    assert {"key": "users", "title": "Users"} in log["masters"] and log["more"] is False
    assert len(admin.get("/api/audit", params={"text": "premium"}).json()["entries"]) == 1
    today = date.today().isoformat()
    assert admin.get("/api/audit", params={"date_to": "2020-01-01"}).json()["entries"] == []
    assert len(admin.get("/api/audit", params={"date_from": today, "master": "cars"}
                         ).json()["entries"]) == 2
    excel = admin.get("/api/audit/export", params={"master": "cars"})
    assert excel.status_code == 200 and "Audit_log" in excel.headers["content-disposition"]
    assert load_workbook(io.BytesIO(excel.content)).active.max_row == 3


def test_bring_across_copies_the_desktop_data_and_keeps_users(tmp_path):
    desktop = tmp_path / "desktop" / "drivenstyle.db"
    desktop.parent.mkdir()
    conn = connect(desktop)
    MastersRepo(conn, "desktop user").save("cars", [RowChange(
        None, dict(make="Tata", model="Punch EV", segment="Compact SUV", active=True))])
    conn.close()
    before = desktop.read_bytes()

    app, settings = make(tmp_path)                 # the web tool, with its first admin
    sign_in(app, ADMIN).post("/api/users", json={"email": STAFF}, headers=H)
    done = bring_across(desktop, settings.db_path)
    assert done["cars"] == 1 and done["users"] == 2
    assert desktop.read_bytes() == before                         # source untouched
    assert len(list(settings.data_dir.glob("*.before-bring-across-*.db"))) == 1

    staff = sign_in(app, STAFF)                                   # still allowed in
    assert staff.get("/api/masters/cars/rows").json()["rows"][0]["model"] == "Punch EV"
    assert sign_in(app, ADMIN).get("/api/audit", params={"master": "cars"}
                                   ).json()["entries"][0]["user"] == "desktop user"

    with pytest.raises(BringAcrossError, match="already has masters"):
        bring_across(desktop, settings.db_path)
    assert bring_across(desktop, settings.db_path, replace=True)["cars"] == 1
    with pytest.raises(BringAcrossError, match="no database"):
        bring_across(tmp_path / "missing.db", settings.db_path)
