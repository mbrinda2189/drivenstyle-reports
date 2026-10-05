"""
v0.20.0 (Brinda, 05-10-2026) - presentation rules asked for by the client:
  * incentives typed on Monthly inputs are a DIRECT cost in the P&L;
  * the automatic indirect costs are a percentage of the product cost only,
    and the percentage is not printed beside the head;
  * labour is shown in separate tables for floor mat and sunfilm;
  * counter sales and cars without a segment are counted under "Others".
"""
from openpyxl import Workbook

from app.reports.data import MonthData, labour_group, segment_group
from app.reports.workbook import pnl_sheet


def test_labour_group_and_segment_group():
    assert labour_group("I20 - PVC Full Floor Mat + Labour Extra") == "Floor mat"
    assert labour_group("Sunfilm - Nano Ceramic - Front (SK)") == "Sunfilm"
    assert labour_group("PPF - Full body") == "Other"
    assert segment_group("") == segment_group("Counter sale (no vehicle)") == "Others"
    assert segment_group("Hatchback") == "Hatchback"


def test_incentives_are_a_direct_cost_in_the_pnl():
    d = MonthData(year=2026, month=9, invoices=[], left_out=[], skipped_files=[],
                  indirect_costs=[("Rent", 1000.0), ("Incentives", 500.0)],
                  threshold=40.0, has_inputs=True)
    assert d.direct_incentives == [("Incentives", 500.0)]
    assert d.other_indirect == [("Rent", 1000.0)]
    ws = Workbook().active
    pnl_sheet(ws, d)
    col = {ws.cell(r, 1).value: (r, ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)}
    product, direct, indirect = col["Product cost"][0], col["Total direct costs"], col["Total indirect costs"]
    # Incentives sits between Product cost and Total direct costs ...
    assert product < col["Incentives"][0] < direct[0] < col["Rent"][0] < indirect[0]
    assert direct[1] == f"=SUM(B{product}:B{col['Incentives'][0]})"
    # ... and the automatic heads are a share of the PRODUCT COST row, with
    # no "(4% of COGS)" beside the name.
    assert col["Breakage / returns / transport"][1] == f"=ROUND(B{product}*0.04,2)"
    assert col["Compliance GST"][1] == f"=ROUND(B{product}*0.03,2)"
    assert not any("COGS" in str(k) for k in col if k and ws.cell(col[k][0], 1).font.sz != 9)
