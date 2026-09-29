"""
test_issues_export.py - Tests for "Export issues" on Scan review (v0.6.2)
=========================================================================
"""

from openpyxl import load_workbook

from app.data.database import connect
from app.data.invoice_export import classify_export
from app.data.invoices_repo import FIRM_GSTIN, InvoicesRepo
from app.data.masters_repo import MastersRepo
from app.reports.issues_export import default_file_name, export_issues, what_we_need
from app.reports.workbook import QTY
from tests.test_invoice_export import extra_rows, rows_0753, rows_226, write_csv


def month_repo(tmp_path):
    f = write_csv(tmp_path / "Invoice.csv", rows_0753() + rows_226() + extra_rows())
    irepo = InvoicesRepo(MastersRepo(connect(":memory:")))
    irepo.store_scan(2026, 9, str(f), classify_export(f, 2026, 9, FIRM_GSTIN).results)
    return irepo


def table_rows(ws, heading: str) -> list[tuple]:
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r[0] == heading)
    return [rows[head]] + [r for r in rows[head + 1:] if r[0] is not None]


def test_open_issues_exported_with_what_we_need(tmp_path):
    irepo = month_repo(tmp_path)
    out = tmp_path / default_file_name(2026, 9)
    assert out.name == "DriveNStyle_Sep-2026_Issues.xlsx"
    n = export_issues(irepo, 2026, 9, out)
    open_issues = [i for i in irepo.issues(2026, 9) if i.status == "open"]
    assert n == len(open_issues) > 0

    wb = load_workbook(out)
    assert wb.sheetnames == ["Issues", "Invoices affected"]
    rows = table_rows(wb["Issues"], "Type")
    assert rows[0] == ("Type", "As printed on the invoice", "Issue", "Invoices",
                       "Invoice numbers", "What we need", "Client's reply")
    assert len(rows) - 1 == n
    sp = [r for r in rows if r[0] == "Salesperson" and r[1] == "Nandha Kumar"]
    assert sp and "contact no and branch" in sp[0][5] and sp[0][6] is None
    totals = [r for r in rows if r[0] == "Totals"]
    assert totals and totals[0][4] == "DNS-301-2627"

    inv = table_rows(wb["Invoices affected"], "Invoice no")
    numbers = {r[0] for r in inv[1:]}
    assert numbers == {no for i in open_issues for no in i.invoices}
    assert "DNS-300-2627" not in numbers            # void: skipped, not an open issue


def test_what_we_need_depends_on_the_reason():
    assert "another branch" in what_we_need("salesperson", "branch")
    assert "who made this sale" in what_we_need("salesperson", "several")
    assert "Product master" in what_we_need("product", "")


def test_quantity_format_has_no_trailing_dot():
    assert QTY == "General"                          # "#,##0.##" showed "9."
