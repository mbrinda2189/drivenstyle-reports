"""
test_reports.py - Tests for monthly inputs, the payments export and the reports
==============================================================================

Invoices are built in code (see test_invoices.py); the workbook is written
to a temporary folder and read back with openpyxl. Formula RESULTS are
checked by opening the file in Excel/LibreOffice, not here; these tests check
the figures written and that every report sheet is present.
"""

from datetime import date

import pytest
from openpyxl import load_workbook

from app.data.inputs_repo import InputsRepo
from app.data.invoices_repo import FileResult, InvoicesRepo
from app.data.masters_repo import RowChange
from app.data.payments_io import PaymentsFileError, read_payments
from app.reports.data import build_month
from app.reports.generate import GenerateError, file_name, generate
from app.sample_data import REPORTS
from tests.test_invoices import inv_0753, inv_226


@pytest.fixture
def world(repo):
    """Masters, a scanned September with two invoices, and inputs."""
    inv, inp = InvoicesRepo(repo), InputsRepo(repo)
    repo.save("incentives", [RowChange(None, dict(name="Underbody", incentive_amount=200,
                                                  bill_value=3500, active=True), date(2026, 4, 1)),
                             RowChange(None, dict(name="Basic Package", incentive_amount=1000,
                                                  bill_value=15000, active=True), date(2026, 4, 1))])
    repo.save("executives", [RowChange(None, dict(name="Edhayan", phone="9000000004",
                                                  branch="Ooty", active=True)),
                             RowChange(None, dict(name="Nandha Kumar", phone="9000000003",
                                                  branch="HO", active=True))])
    repo.save("cars", [RowChange(None, dict(make="Hyundai", model="i20", segment="Hatchback", active=True)),
                       RowChange(None, dict(make="Tata", model="Punch EV", segment="Compact SUV", active=True))])
    items = [("I20 - PVC Full Floor Mat + Labour Extra", "Product", 1800, True, 150, ""),
             ("Underbody Coating - 5 Seater", "Service", 900, False, 0, "Underbody"),
             ("Silencer Coating - All Cars", "Service", 300, False, 0, ""),
             ("Sunfilm - Nano Ceramic - Front (SK)", "Service", 2400, True, 300, ""),
             ("Sunfilm - Nano Ceramic - Side and Rear (SK)", "Service", 3100, True, 500, ""),
             ("Horn", "Product", 1100, False, 0, ""),
             ("Punch.EV Mudflap - Techno", "Product", 300, False, 0, ""),
             ("PUNCH - PVC Full Floor Mat + Labour Extra", "Product", 1900, True, 150, "Basic Package"),
             ("Punch EV Number Plate Frame", "Product", 200, False, 0, "")]
    repo.save("products", [RowChange(None, dict(
        sku="", name=n, hsn_sac="", category=c, incentive_group=g, selling_price=0,
        cost_price=cp, has_labour=l, labour_charge=lc, active=True), date(2026, 4, 1))
        for n, c, cp, l, lc, g in items])
    inv.store_scan(2026, 9, "folder", [FileResult(i.file_name, "read", "", i)
                                       for i in (inv_0753(), inv_226())])
    inp.save_costs(2026, 9, [("Rent", 45000), ("Salaries", 120000)])
    return repo, inv, inp


# --- monthly inputs -----------------------------------------------------------
def test_costs_are_saved_logged_and_copied(world):
    repo, _, inp = world
    assert inp.costs(2026, 9) == [("Rent", 45000.0), ("Salaries", 120000.0)]
    inp.save_costs(2026, 9, [("Rent", 50000), ("Power", 9000)])
    log = {(e["action"], e["field"]) for e in repo.audit_entries(master="inputs")}
    assert {("Edited", "Rent"), ("Deleted", "Salaries"), ("Added", "Power")} <= log
    label, heads = inp.previous_costs(2026, 10)
    assert label == "September 2026" and heads == [("Rent", 50000.0), ("Power", 9000.0)]


def test_cost_heads_must_be_unique_and_not_negative(world):
    _, _, inp = world
    with pytest.raises(ValueError):
        inp.save_costs(2026, 9, [("Rent", 1), ("rent", 2)])
    with pytest.raises(ValueError):
        inp.save_costs(2026, 9, [("Rent", -1)])


def test_threshold_default_and_saved(world):
    _, _, inp = world
    assert inp.threshold(2026, 8) == 40.0
    inp.save_threshold(2026, 9, 45)
    assert inp.threshold(2026, 9) == 45.0


# --- payments export -------------------------------------------------------------
def test_read_payments_csv(tmp_path):
    f = tmp_path / "Customer_Payment.csv"
    f.write_text("Payment Number,Mode,Amount,Date,Deposit To,Amount Applied to Invoice,"
                 "Invoice Number\n1,UPI,25000,2026-09-01,HDFC,25000,DNS-159-2627\n"
                 "2,Bank Transfer,50000,2026-09-01,HDFC,50000,DNS-159-2627\n", encoding="utf-8")
    p = read_payments(f)
    assert [(x.invoice_no, x.mode, x.amount) for x in p] == [
        ("DNS-159-2627", "UPI", 25000.0), ("DNS-159-2627", "Bank Transfer", 50000.0)]
    assert p[0].date == date(2026, 9, 1)


def test_read_payments_refuses_other_files(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("a,b\n1,2\n")
    with pytest.raises(PaymentsFileError):
        read_payments(f)


# --- figures -----------------------------------------------------------------------
def test_month_figures(world):
    repo, inv, inp = world
    d = build_month(repo, inv, inp, 2026, 9)
    assert len(d.invoices) == 2 and not d.left_out
    # v0.6.3: the Rs. 1 labour lines (1 on 0753, 3 on 226) are ignored -
    # not in sales; their value is reported separately.
    assert d.marker_lines == 4 and 0 < d.marker_value < 4
    assert d.sales == pytest.approx(7033.75 + 20000 - d.marker_value, abs=0.05)
    i226 = next(i for i in d.invoices if i.invoice_no == "DNS-226-2627")
    # cost: 2400 + 3100 + 1100 + 300 + 1900 + 200; labour: 300 + 500 + 150
    assert i226.cost == 9000 and i226.labour == 950
    assert i226.markers == 3 and not any(l.category == "Labour line" for l in d.lines)
    assert i226.labour == 950                                  # from the products
    # v0.8.0: a product linked to a "... Package" incentive group is no
    # longer a package sale by itself (see tests/test_packages.py)
    assert not any(l.is_package for l in d.lines)


def test_open_issue_leaves_invoice_out(world):
    repo, inv, inp = world
    first = [p for p in repo.list_rows("products") if p["name"] == "Horn"][0]
    repo.delete("products", [first["id"]])                      # Horn now unknown
    d = build_month(repo, inv, inp, 2026, 9)
    assert [x.invoice_no for x in d.left_out] == ["DNS-226-2627"]
    assert "Horn" in d.left_out[0].reasons[0]


def test_rates_on_the_invoice_date_are_used(world):
    repo, inv, inp = world
    p = [x for x in repo.list_rows("products") if x["name"] == "Horn"][0]
    repo.save("products", [RowChange(p["id"], dict(p, cost_price=5000.0), date(2026, 10, 1))])
    d = build_month(repo, inv, inp, 2026, 9)                    # September: old cost
    horn = [l for l in d.lines if l.product == "Horn"][0]
    assert horn.cost == 1100


# --- workbook ----------------------------------------------------------------------
def test_generate_writes_every_sheet(world, tmp_path):
    repo, inv, inp = world
    pay = tmp_path / "p.csv"
    pay.write_text("Mode,Amount Applied to Invoice,Invoice Number,Deposit To\n"
                   "UPI,20000,DNS-226-2627,HDFC\nCash,5000,DNS26-GST-0753,Undeposited\n",
                   encoding="utf-8")
    r = generate(repo, inv, inp, 2026, 9, tmp_path, REPORTS, str(pay))
    assert file_name(2026, 9) == "DriveNStyle_Sep-2026_Reports.xlsx"
    # v0.8.0: the file name carries the date and time it was generated
    import re
    assert re.fullmatch(r"DriveNStyle_Sep-2026_Reports_\d\d-\d\d-\d{4}_\d{4}\.xlsx",
                        r.path.name)
    wb = load_workbook(r.path)
    # v0.9.0: + the executive summary, right after the cover
    assert len(wb.sheetnames) == 15 and wb.sheetnames[:2] == ["Cover", "Summary"]
    ws = wb["11 Payment modes"]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Cash" in texts and "UPI" in texts
    formulas = [c.value for row in wb["1 Invoice profitability"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUM(" in f for f in formulas)
    assert inp.runs()[0]["invoices"] == 2


def test_generate_only_ticked_reports(world, tmp_path):
    repo, inv, inp = world
    r = generate(repo, inv, inp, 2026, 9, tmp_path, ["Profit & loss"])
    assert load_workbook(r.path).sheetnames == [
        "Cover", "Summary", "12 Profit & loss", "Not included"]


def test_generate_needs_a_scanned_month(world, tmp_path):
    repo, inv, inp = world
    with pytest.raises(GenerateError):
        generate(repo, inv, inp, 2026, 8, tmp_path, REPORTS)